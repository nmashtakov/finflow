import csv
import io
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.files.base import ContentFile
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Case, Count, DecimalField, ExpressionWrapper, F, IntegerField, Prefetch, Q, Sum, Value, When
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.text import get_valid_filename
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from .access import crm_access_required, get_current_organization
from .forms import (
    CounterpartyForm,
    DeliveryMethodForm,
    OrderForm,
    OrderItemFormSet,
    OrganizationOnboardingForm,
    PurchaseCSVUploadForm,
    PurchaseForm,
    PurchaseItemFormSet,
    ProductBulkUploadForm,
    ProductChannelOfferForm,
    ProductForm,
    SalesChannelForm,
    WarehouseForm,
)
from .models import (
    Counterparty,
    DeliveryMethod,
    Membership,
    Order,
    OrderItem,
    Product,
    ProductChannelOffer,
    Purchase,
    SalesChannel,
    StockLot,
    StockReservation,
    Warehouse,
)
from .services import (
    PURCHASE_TEMPLATE_HEADERS,
    PRODUCT_IMPORT_HEADER_ALIASES,
    PRODUCT_IMPORT_TEMPLATE_HEADERS,
    get_order_detail_summary,
    initialize_organization_defaults,
    release_order_reservations,
    replace_order_items_and_reservations,
    replace_purchase_items_and_stock,
)


PRODUCTS_PER_PAGE = 25
ORDERS_PER_PAGE = 25
PRODUCT_IMPORT_SUMMARY_SESSION_KEY = 'crm_product_import_summary'
PRODUCT_SORT_OPTIONS = {
    'name_asc': {
        'label': 'Название: А-Я',
        'ordering': ['-is_active', 'name', 'sku', 'id'],
    },
    'name_desc': {
        'label': 'Название: Я-А',
        'ordering': ['-is_active', '-name', 'sku', 'id'],
    },
    'sku_asc': {
        'label': 'SKU: по возрастанию',
        'ordering': ['-is_active', 'sku', 'id'],
    },
    'sku_desc': {
        'label': 'SKU: по убыванию',
        'ordering': ['-is_active', '-sku', 'id'],
    },
    'stock_desc': {
        'label': 'Остатки: больше сверху',
        'ordering': ['-is_active', '-total_stock_quantity', 'name', 'sku', 'id'],
    },
    'stock_asc': {
        'label': 'Остатки: меньше сверху',
        'ordering': ['-is_active', 'total_stock_quantity', 'name', 'sku', 'id'],
    },
    'cost_desc': {
        'label': 'Себестоимость: больше сверху',
        'ordering': ['-is_active', '-calculated_unit_cost', 'name', 'sku', 'id'],
    },
    'cost_asc': {
        'label': 'Себестоимость: меньше сверху',
        'ordering': ['-is_active', 'calculated_unit_cost', 'name', 'sku', 'id'],
    },
}
ORDER_STATUS_BADGES = {
    Order.Status.NEW: ('warning', 'Новый'),
    Order.Status.AWAITING_SHIPMENT: ('primary', 'Ожидает отправки / самовывоз'),
    Order.Status.IN_DELIVERY: ('info', 'В доставке'),
    Order.Status.PICKUP_POINT: ('secondary', 'В ПВЗ'),
    Order.Status.COMPLETED: ('success', 'Завершен'),
    Order.Status.CANCELLED: ('dark', 'Отменен'),
    Order.Status.RETURNED: ('danger', 'Возврат'),
}


@crm_access_required
def dashboard_view(request):
    organization = request.crm_organization
    return render(
        request,
        'crm/dashboard.html',
        {
            'organization': organization,
            'dashboard_cards': [
                ('Новые заказы', '—'),
                ('Продажи сегодня', '—'),
                ('Проблемы', '—'),
                ('Остатки', '—'),
            ],
        },
    )


def _get_organization_object_or_404(request, model, object_id):
    return get_object_or_404(
        model,
        pk=object_id,
        organization=request.crm_organization,
    )


def _render_directory_form(
    request,
    form,
    page_title,
    submit_label,
    cancel_url,
):
    return render(
        request,
        'crm/directory_form.html',
        {
            'organization': request.crm_organization,
            'form': form,
            'page_title': page_title,
            'submit_label': submit_label,
            'cancel_url': cancel_url,
        },
    )


def _get_listing_channels(organization):
    return SalesChannel.objects.filter(
        organization=organization,
        is_active=True,
        is_listing_channel=True,
    ).order_by('name', 'id')


def _get_purchase_products_queryset(organization, purchase=None):
    used_product_ids = []
    if purchase is not None and purchase.pk:
        used_product_ids = purchase.items.values_list('product_id', flat=True)
    return Product.objects.filter(
        organization=organization,
    ).filter(
        Q(is_active=True) | Q(id__in=used_product_ids)
    ).order_by('name', 'sku', 'id')


def _get_purchase_warehouses_queryset(organization, purchase=None):
    used_warehouse_ids = []
    if purchase is not None and purchase.pk:
        used_warehouse_ids = purchase.items.values_list('warehouse_id', flat=True)
    return Warehouse.objects.filter(
        organization=organization,
    ).filter(
        Q(is_active=True) | Q(id__in=used_warehouse_ids)
    ).order_by('name', 'id')


def _build_purchase_item_formset(request, purchase=None, data=None, initial=None):
    return PurchaseItemFormSet(
        data=data,
        initial=initial,
        prefix='items',
        form_kwargs={
            'product_queryset': _get_purchase_products_queryset(request.crm_organization, purchase=purchase),
        },
    )


def _get_purchase_item_initial_rows(purchase):
    if purchase is None or not purchase.pk:
        return None

    return [
        {
            'product': item.product_id,
            'quantity': item.quantity,
            'unit_cost': item.unit_cost,
            'extra_cost_total': item.extra_cost_total,
            'comment': item.comment,
        }
        for item in purchase.items.select_related('product', 'warehouse').order_by('id')
    ]


def _get_purchase_detail_summary(purchase):
    items = list(
        purchase.items.select_related('product', 'warehouse')
        .order_by('id')
    )
    total_quantity = sum(item.quantity for item in items)
    total_amount = sum((item.line_total for item in items), Decimal('0.00'))
    return {
        'items': items,
        'line_count': len(items),
        'total_quantity': total_quantity,
        'total_amount': total_amount,
    }


def _get_order_products_queryset(organization, order=None):
    used_product_ids = []
    if order is not None and order.pk:
        used_product_ids = order.items.values_list('product_id', flat=True)
    return Product.objects.filter(
        organization=organization,
    ).filter(
        Q(is_active=True) | Q(id__in=used_product_ids)
    ).order_by('name', 'sku', 'id')


def _get_order_sales_channels_queryset(organization, order=None):
    used_channel_id = getattr(order, 'sales_channel_id', None)
    return SalesChannel.objects.filter(
        organization=organization,
    ).filter(
        Q(is_active=True) | Q(id=used_channel_id)
    ).order_by('name', 'id')


def _get_order_delivery_methods_queryset(organization, order=None):
    used_method_id = getattr(order, 'delivery_method_id', None)
    return DeliveryMethod.objects.filter(
        organization=organization,
    ).filter(
        Q(is_active=True) | Q(id=used_method_id)
    ).order_by('sort_order', 'name', 'id')


def _build_order_item_formset(request, order=None, data=None, initial=None):
    return OrderItemFormSet(
        data=data,
        initial=initial,
        prefix='items',
        form_kwargs={
            'product_queryset': _get_order_products_queryset(request.crm_organization, order=order),
        },
    )


def _get_order_item_initial_rows(order):
    if order is None or not order.pk:
        return None

    return [
        {
            'product': item.product_id,
            'quantity': item.quantity,
            'sale_price': item.sale_price,
            'comment': item.comment,
        }
        for item in order.items.select_related('product').order_by('id')
    ]


def _get_order_sales_channel_price_map(organization):
    offer_map = {}
    offers = ProductChannelOffer.objects.filter(
        product__organization=organization,
        channel__organization=organization,
        is_enabled=True,
        price__isnull=False,
    ).select_related('product', 'channel')

    for offer in offers:
        offer_map.setdefault(str(offer.product_id), {})[str(offer.channel_id)] = str(offer.price)
    return offer_map


def _get_order_product_picker_products(organization, order=None):
    return (
        _get_order_products_queryset(organization, order=order)
        .annotate(
            total_stock_quantity=Coalesce(
                Sum(
                    'stock_lots__remaining_quantity',
                    filter=Q(stock_lots__remaining_quantity__gt=0),
                    output_field=IntegerField(),
                ),
                Value(0),
                output_field=IntegerField(),
            )
        )
    )


def _apply_order_summary(order):
    summary = get_order_detail_summary(order)
    order.summary = summary
    order.reservation_status_code = summary['reservation_status_code']
    order.reservation_status_label = summary['reservation_status_label']
    order.missing_quantity_total = summary['missing_quantity_total']
    order.total_sale_amount_value = summary['total_sale_amount']
    status_badge_class, status_badge_label = ORDER_STATUS_BADGES.get(
        order.status,
        ('secondary', order.get_status_display()),
    )
    order.status_badge_class = status_badge_class
    order.status_badge_label = status_badge_label
    return summary


def _resolve_purchase_import_warehouse(raw_warehouse_id, organization):
    warehouse_id = (raw_warehouse_id or '').strip()
    if not warehouse_id:
        raise ValueError('Перед загрузкой CSV выберите склад закупки.')

    warehouse = Warehouse.objects.filter(
        organization=organization,
        is_active=True,
        id=warehouse_id,
    ).first()
    if warehouse is None:
        raise ValueError('Выбранный склад не найден в текущей организации.')
    return warehouse


def _normalize_csv_numeric_value(raw_value):
    return (
        str(raw_value or '')
        .strip()
        .replace('\xa0', '')
        .replace(' ', '')
        .replace(',', '.')
    )


def _build_purchase_csv_preview(rows, organization, warehouse):
    product_map = Product.objects.filter(
        organization=organization,
        id__in=[row['product'] for row in rows],
    ).in_bulk()

    return {
        'warehouse_name': warehouse.name,
        'total_rows': len(rows),
        'rows': [
            {
                'product_label': f"{product_map[row['product']].sku} — {product_map[row['product']].name}",
                'quantity': row['quantity'],
                'unit_cost': row['unit_cost'],
                'extra_cost_total': row['extra_cost_total'],
                'comment': row['comment'],
            }
            for row in rows[:5]
            if row['product'] in product_map
        ],
    }


def _parse_purchase_csv(file, organization):
    text_stream = io.StringIO(file.read().decode('utf-8-sig'))
    reader = csv.DictReader(text_stream)
    if reader.fieldnames is None:
        raise ValueError('CSV-файл пустой.')

    normalized_headers = [header.strip() for header in reader.fieldnames]
    missing_headers = [header for header in PURCHASE_TEMPLATE_HEADERS if header not in normalized_headers]
    if missing_headers:
        raise ValueError(f'В CSV не хватает колонок: {", ".join(missing_headers)}.')

    products = {
        product.sku: product
        for product in Product.objects.filter(organization=organization, is_active=True)
    }

    rows = []
    for line_number, row in enumerate(reader, start=2):
        sku = (row.get('sku') or '').strip()
        quantity_raw = (row.get('quantity') or '').strip()
        unit_cost_raw = (row.get('unit_cost') or '').strip()
        extra_cost_raw = (row.get('extra_cost_total') or '').strip()
        comment = (row.get('comment') or '').strip()

        if not any([sku, quantity_raw, unit_cost_raw, extra_cost_raw, comment]):
            continue

        product = products.get(sku)
        if product is None:
            raise ValueError(f'Строка {line_number}: товар с SKU "{sku}" не найден в текущей организации.')

        try:
            quantity_decimal = Decimal(_normalize_csv_numeric_value(quantity_raw))
        except (InvalidOperation, TypeError):
            raise ValueError(f'Строка {line_number}: количество должно быть целым числом.')
        if quantity_decimal != quantity_decimal.to_integral_value():
            raise ValueError(f'Строка {line_number}: количество должно быть целым числом.')
        quantity = int(quantity_decimal)
        if quantity <= 0:
            raise ValueError(f'Строка {line_number}: количество должно быть больше нуля.')

        try:
            unit_cost = Decimal(_normalize_csv_numeric_value(unit_cost_raw))
        except (InvalidOperation, TypeError):
            raise ValueError(f'Строка {line_number}: цена закупки должна быть числом.')
        if unit_cost < 0:
            raise ValueError(f'Строка {line_number}: цена закупки не может быть отрицательной.')

        try:
            extra_cost_total = Decimal(_normalize_csv_numeric_value(extra_cost_raw or '0'))
        except (InvalidOperation, TypeError):
            raise ValueError(f'Строка {line_number}: дополнительные расходы должны быть числом.')
        if extra_cost_total < 0:
            raise ValueError(f'Строка {line_number}: дополнительные расходы не могут быть отрицательными.')

        rows.append(
            {
                'product': product.id,
                'quantity': quantity,
                'unit_cost': unit_cost,
                'extra_cost_total': extra_cost_total,
                'comment': comment,
            }
        )

    if not rows:
        raise ValueError('В CSV нет строк для загрузки.')

    return rows


def _parse_import_boolean(value, *, default=True):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)

    normalized_value = str(value).strip().lower()
    if normalized_value == '':
        return default
    if normalized_value in {'1', 'true', 'yes', 'y', 'on', 'да', 'активен', 'активный'}:
        return True
    if normalized_value in {'0', 'false', 'no', 'n', 'off', 'нет', 'архив', 'в архиве', 'неактивен'}:
        return False
    raise ValueError(f'Не удалось распознать булево значение "{value}".')


def _normalize_import_header(value):
    return ' '.join(
        str(value or '')
        .strip()
        .lower()
        .replace('_', ' ')
        .replace('-', ' ')
        .replace('\n', ' ')
        .split()
    )


def _stringify_excel_value(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _resolve_product_import_columns(worksheet):
    header_lookup = {}
    for index, cell in enumerate(worksheet[1], start=1):
        normalized_header = _normalize_import_header(cell.value)
        if normalized_header and normalized_header not in header_lookup:
            header_lookup[normalized_header] = index

    column_map = {}
    for field_name, aliases in PRODUCT_IMPORT_HEADER_ALIASES.items():
        for alias in aliases:
            normalized_alias = _normalize_import_header(alias)
            if normalized_alias in header_lookup:
                column_map[field_name] = header_lookup[normalized_alias]
                break

    missing_fields = [field_name for field_name in ('sku', 'name') if field_name not in column_map]
    if missing_fields:
        raise ValueError(
            'В Excel не хватает обязательных колонок для SKU и названия товара. '
            'Поддерживаются заголовки вроде SKU / Номер и Name / Название / Набор.'
        )

    return column_map


def _extract_excel_photo_candidates(worksheet):
    photo_candidates = []
    for image in getattr(worksheet, '_images', []):
        anchor = getattr(image, 'anchor', None)
        marker = getattr(anchor, '_from', None)
        if marker is None:
            continue

        row_number = marker.row + 1
        column_number = marker.col + 1
        image_format = (getattr(image, 'format', None) or 'png').lower()
        if image_format == 'jpeg':
            image_format = 'jpg'
        photo_candidates.append(
            {
                'row_number': row_number,
                'column_number': column_number,
                'content': image._data(),
                'extension': image_format,
            }
        )

    return sorted(photo_candidates, key=lambda item: (item['row_number'], item['column_number']))


def _normalize_zip_target(base_path, target):
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_path), target))


def _extract_rich_data_photo_candidates(file_bytes):
    namespaces = {
        'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
        'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'rvrel': 'http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel',
    }

    photo_candidates = []
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            workbook_root = ET.fromstring(archive.read('xl/workbook.xml'))
            workbook_rels_root = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
            workbook_view = workbook_root.find('main:bookViews/main:workbookView', namespaces)
            active_tab = int(workbook_view.attrib.get('activeTab', '0')) if workbook_view is not None else 0
            sheets = workbook_root.findall('main:sheets/main:sheet', namespaces)
            if not sheets:
                return []

            active_tab = min(active_tab, len(sheets) - 1)
            sheet_rel_id = sheets[active_tab].attrib.get(f'{{{namespaces["r"]}}}id')
            workbook_rel_targets = {
                rel.attrib['Id']: rel.attrib['Target']
                for rel in workbook_rels_root.findall('rel:Relationship', namespaces)
            }
            sheet_target = workbook_rel_targets.get(sheet_rel_id)
            if not sheet_target:
                return []

            sheet_path = _normalize_zip_target('xl/workbook.xml', sheet_target)
            sheet_root = ET.fromstring(archive.read(sheet_path))
            metadata_root = ET.fromstring(archive.read('xl/metadata.xml'))
            rich_value_root = ET.fromstring(archive.read('xl/richData/rdrichvalue.xml'))
            rich_value_rel_root = ET.fromstring(archive.read('xl/richData/richValueRel.xml'))
            rich_value_rel_rels_root = ET.fromstring(archive.read('xl/richData/_rels/richValueRel.xml.rels'))

            value_metadata = metadata_root.find('main:valueMetadata', namespaces)
            if value_metadata is None:
                return []

            vm_to_rich_value_index = {}
            for index, bk in enumerate(value_metadata.findall('main:bk', namespaces), start=1):
                rc = bk.find('main:rc', namespaces)
                if rc is None:
                    continue
                vm_to_rich_value_index[index] = int(rc.attrib.get('v', '0'))

            rich_values = rich_value_root.findall('{http://schemas.microsoft.com/office/spreadsheetml/2017/richdata}rv')
            relation_ids_in_order = [
                rel.attrib.get(f'{{{namespaces["r"]}}}id')
                for rel in rich_value_rel_root.findall('{http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel}rel')
            ]
            relation_targets = {
                rel.attrib['Id']: _normalize_zip_target('xl/richData/richValueRel.xml', rel.attrib['Target'])
                for rel in rich_value_rel_rels_root.findall('rel:Relationship', namespaces)
            }

            for cell in sheet_root.findall('.//main:c[@vm]', namespaces):
                vm_index = int(cell.attrib['vm'])
                rich_value_index = vm_to_rich_value_index.get(vm_index)
                if rich_value_index is None or rich_value_index >= len(rich_values):
                    continue

                rich_value = rich_values[rich_value_index]
                rich_value_values = rich_value.findall('{http://schemas.microsoft.com/office/spreadsheetml/2017/richdata}v')
                if not rich_value_values:
                    continue

                relation_index = int((rich_value_values[0].text or '0').strip())
                if relation_index >= len(relation_ids_in_order):
                    continue

                relation_id = relation_ids_in_order[relation_index]
                target_path = relation_targets.get(relation_id)
                if not relation_id or not target_path:
                    continue

                row_number, column_number = coordinate_to_tuple(cell.attrib['r'])
                image_bytes = archive.read(target_path)
                extension = posixpath.splitext(target_path)[1].lstrip('.').lower() or 'png'
                if extension == 'jpeg':
                    extension = 'jpg'

                photo_candidates.append(
                    {
                        'row_number': row_number,
                        'column_number': column_number,
                        'content': image_bytes,
                        'extension': extension,
                    }
                )
    except KeyError:
        return []

    return sorted(photo_candidates, key=lambda item: (item['row_number'], item['column_number']))


def _pick_photo_candidate_for_row(
    row_number,
    photo_candidates,
    used_candidate_indexes,
    preferred_column=None,
    max_row_distance=0,
):
    best_match = None
    for index, candidate in enumerate(photo_candidates):
        if index in used_candidate_indexes:
            continue

        row_distance = abs(candidate['row_number'] - row_number)
        if row_distance > max_row_distance:
            continue

        column_distance = abs(candidate['column_number'] - preferred_column) if preferred_column else 0
        ranking = (row_distance, column_distance, candidate['column_number'])
        if best_match is None or ranking < best_match[0]:
            best_match = (
                ranking,
                index,
                {
                    'content': candidate['content'],
                    'extension': candidate['extension'],
                },
            )

    if best_match is None:
        return None

    used_candidate_indexes.add(best_match[1])
    return best_match[2]


def _attach_photos_to_import_rows(rows, photo_candidates, preferred_column=None):
    used_candidate_indexes = set()

    for row in rows:
        photo_data = _pick_photo_candidate_for_row(
            row['row_number'],
            photo_candidates,
            used_candidate_indexes,
            preferred_column=preferred_column,
            max_row_distance=0,
        )
        if photo_data is not None:
            row['photo_name'] = _build_product_photo_name(
                row['sku'],
                row['row_number'],
                photo_data['extension'],
            )
            row['photo_content'] = photo_data['content']

    for row in rows:
        if row['photo_content'] is not None:
            continue
        photo_data = _pick_photo_candidate_for_row(
            row['row_number'],
            photo_candidates,
            used_candidate_indexes,
            preferred_column=preferred_column,
            max_row_distance=2,
        )
        if photo_data is not None:
            row['photo_name'] = _build_product_photo_name(
                row['sku'],
                row['row_number'],
                photo_data['extension'],
            )
            row['photo_content'] = photo_data['content']


def _build_product_photo_name(sku, row_number, extension):
    file_name = get_valid_filename(f'{sku}-{row_number}.{extension}')
    return file_name or f'product-{row_number}.{extension}'


def _build_product_import_summary(skipped_rows):
    grouped_by_reason = {}
    for row in skipped_rows:
        reason_group = grouped_by_reason.setdefault(
            row['reason'],
            {
                'reason': row['reason'],
                'count': 0,
                'items_by_sku': {},
            },
        )
        reason_group['count'] += 1
        reason_group['items_by_sku'].setdefault(row['sku'], []).append(row['row_number'])

    groups = []
    for group in grouped_by_reason.values():
        items = [
            {
                'sku': sku,
                'row_numbers': row_numbers,
                'row_numbers_display': ', '.join(str(row_number) for row_number in row_numbers),
            }
            for sku, row_numbers in group['items_by_sku'].items()
        ]
        groups.append(
            {
                'reason': group['reason'],
                'count': group['count'],
                'items': items,
            }
        )

    return {
        'total_count': len(skipped_rows),
        'groups': groups,
    }


def _pop_product_import_summary(request):
    return request.session.pop(PRODUCT_IMPORT_SUMMARY_SESSION_KEY, None)


def _parse_product_import_xlsx(xlsx_file):
    xlsx_file.seek(0)
    file_bytes = xlsx_file.read()
    workbook_stream = io.BytesIO(file_bytes)
    try:
        workbook = load_workbook(workbook_stream, data_only=True)
    except Exception as error:
        raise ValueError('Не удалось прочитать Excel-файл. Используйте .xlsx или .xlsm.') from error

    worksheet = workbook.active
    if worksheet.max_row < 2:
        raise ValueError('Excel-файл пустой.')

    column_map = _resolve_product_import_columns(worksheet)
    photo_candidates = _extract_excel_photo_candidates(worksheet)
    photo_candidates.extend(_extract_rich_data_photo_candidates(file_bytes))
    photo_candidates = sorted(photo_candidates, key=lambda item: (item['row_number'], item['column_number']))
    photo_column = column_map.get('photo')
    rows = []

    for row_number in range(2, worksheet.max_row + 1):
        sku = _stringify_excel_value(worksheet.cell(row=row_number, column=column_map['sku']).value)
        name = _stringify_excel_value(worksheet.cell(row=row_number, column=column_map['name']).value)
        comment = _stringify_excel_value(
            worksheet.cell(row=row_number, column=column_map['comment']).value
        ) if 'comment' in column_map else ''
        is_active_raw = (
            worksheet.cell(row=row_number, column=column_map['is_active']).value
            if 'is_active' in column_map else None
        )

        row_has_input = any([sku, name, comment, is_active_raw not in (None, '')])
        if not row_has_input:
            continue

        if not sku:
            raise ValueError(f'Строка {row_number}: SKU / артикул обязателен.')
        if not name:
            raise ValueError(f'Строка {row_number}: название товара обязательно.')

        try:
            is_active = _parse_import_boolean(is_active_raw, default=True)
        except ValueError as error:
            raise ValueError(f'Строка {row_number}: {error}')

        rows.append(
            {
                'sku': sku,
                'name': name,
                'comment': comment,
                'is_active': is_active,
                'row_number': row_number,
                'photo_name': None,
                'photo_content': None,
            }
        )

    if not rows:
        raise ValueError('В Excel-файле нет строк для загрузки.')

    _attach_photos_to_import_rows(rows, photo_candidates, preferred_column=photo_column)
    return rows


def _create_products_from_rows(organization, rows):
    existing_skus = set(
        Product.objects.filter(
            organization=organization,
            sku__in=[row['sku'] for row in rows],
        ).values_list('sku', flat=True)
    )
    seen_batch_skus = set()
    created_count = 0
    skipped_rows = []

    for row in rows:
        sku = row['sku']
        if sku in seen_batch_skus:
            skipped_rows.append(
                {
                    'sku': sku,
                    'row_number': row['row_number'],
                    'reason': 'дублируется в файле',
                }
            )
            continue

        seen_batch_skus.add(sku)
        if sku in existing_skus:
            skipped_rows.append(
                {
                    'sku': sku,
                    'row_number': row['row_number'],
                    'reason': 'уже существует в CRM',
                }
            )
            continue

        product = Product(
            organization=organization,
            sku=sku,
            name=row['name'],
            comment=row['comment'],
            is_active=row['is_active'],
        )

        if row['photo_content'] is not None:
            product.photo.save(
                row['photo_name'],
                ContentFile(row['photo_content']),
                save=False,
            )

        product.save()
        created_count += 1

    return created_count, skipped_rows


def _get_product_stock_lots(product):
    if product is None or not product.pk:
        return []

    return list(
        StockLot.objects.filter(
            organization=product.organization,
            product=product,
            remaining_quantity__gt=0,
        )
        .select_related('warehouse', 'purchase_item__purchase')
        .order_by('warehouse__name', 'purchase_item__purchase__purchase_date', 'id')
    )


def _save_product_channel_offers(product, offer_form):
    for payload in offer_form.build_offer_payloads():
        existing_offer = payload['existing_offer']
        if existing_offer:
            existing_offer.is_enabled = payload['is_enabled']
            existing_offer.price = payload['price']
            existing_offer.comment = payload['comment']
            existing_offer.save(update_fields=['is_enabled', 'price', 'comment', 'updated_at'])
            continue

        ProductChannelOffer.objects.create(
            product=product,
            channel=payload['channel'],
            is_enabled=payload['is_enabled'],
            price=payload['price'],
            comment=payload['comment'],
        )


def _render_product_form(
    request,
    form,
    offer_form,
    page_title,
    submit_label,
    stock_lots=None,
):
    stock_lots = list(stock_lots or [])
    return render(
        request,
        'crm/product_form.html',
        {
            'organization': request.crm_organization,
            'form': form,
            'offer_form': offer_form,
            'page_title': page_title,
            'submit_label': submit_label,
            'stock_lots': stock_lots,
            'total_stock_quantity': sum(stock_lot.remaining_quantity for stock_lot in stock_lots),
        },
    )


def _render_purchase_form(
    request,
    form,
    item_formset,
    upload_form,
    page_title,
    submit_label,
    purchase=None,
    csv_preview=None,
):
    purchase = purchase if purchase and purchase.pk else None
    detail_summary = _get_purchase_detail_summary(purchase) if purchase else None
    return render(
        request,
        'crm/purchase_form.html',
        {
            'organization': request.crm_organization,
            'form': form,
            'item_formset': item_formset,
            'upload_form': upload_form,
            'purchase': purchase,
            'purchase_detail_summary': detail_summary,
            'page_title': page_title,
            'submit_label': submit_label,
            'csv_preview': csv_preview,
        },
    )


def _render_order_form(
    request,
    form,
    item_formset,
    page_title,
    submit_label,
    order=None,
    order_summary=None,
):
    order = order if order and order.pk else None
    order_summary = order_summary or (_apply_order_summary(order) if order else None)
    return render(
        request,
        'crm/order_form.html',
        {
            'organization': request.crm_organization,
            'form': form,
            'item_formset': item_formset,
            'order': order,
            'order_summary': order_summary,
            'page_title': page_title,
            'submit_label': submit_label,
            'order_status_badges': ORDER_STATUS_BADGES,
            'product_channel_price_map': _get_order_sales_channel_price_map(request.crm_organization),
            'product_picker_products': _get_order_product_picker_products(
                request.crm_organization,
                order=order,
            ),
        },
    )


@crm_access_required
def settings_view(request):
    organization = request.crm_organization
    sales_channels = SalesChannel.objects.filter(organization=organization).order_by('is_active', 'name', 'id')
    warehouses = Warehouse.objects.filter(organization=organization).order_by('is_active', 'name', 'id')
    counterparties = Counterparty.objects.filter(organization=organization).order_by('is_active', 'name', 'id')
    delivery_methods = DeliveryMethod.objects.filter(organization=organization).order_by('is_active', 'sort_order', 'name', 'id')

    return render(
        request,
        'crm/settings.html',
        {
            'organization': organization,
            'active_sales_channels': sales_channels.filter(is_active=True),
            'archived_sales_channels': sales_channels.filter(is_active=False),
            'active_warehouses': warehouses.filter(is_active=True),
            'archived_warehouses': warehouses.filter(is_active=False),
            'active_counterparties': counterparties.filter(is_active=True),
            'archived_counterparties': counterparties.filter(is_active=False),
            'active_delivery_methods': delivery_methods.filter(is_active=True),
            'archived_delivery_methods': delivery_methods.filter(is_active=False),
        },
    )


@crm_access_required
def sales_channel_create_view(request):
    if request.method == 'POST':
        form = SalesChannelForm(request.POST)
        if form.is_valid():
            sales_channel = form.save(commit=False)
            sales_channel.organization = request.crm_organization
            sales_channel.save()
            messages.success(request, 'Канал продаж создан.')
            return redirect('crm:settings')
    else:
        form = SalesChannelForm()

    return _render_directory_form(
        request,
        form,
        page_title='Новый канал продаж',
        submit_label='Создать канал',
        cancel_url='crm:settings',
    )


@crm_access_required
def sales_channel_edit_view(request, channel_id):
    sales_channel = _get_organization_object_or_404(request, SalesChannel, channel_id)

    if request.method == 'POST':
        form = SalesChannelForm(request.POST, instance=sales_channel)
        if form.is_valid():
            form.save()
            messages.success(request, 'Канал продаж обновлен.')
            return redirect('crm:settings')
    else:
        form = SalesChannelForm(instance=sales_channel)

    return _render_directory_form(
        request,
        form,
        page_title=f'Редактирование канала продаж: {sales_channel.name}',
        submit_label='Сохранить',
        cancel_url='crm:settings',
    )


@crm_access_required
@require_POST
def sales_channel_archive_view(request, channel_id):
    sales_channel = _get_organization_object_or_404(request, SalesChannel, channel_id)
    sales_channel.is_active = False
    sales_channel.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Канал продаж отправлен в архив.')
    return redirect('crm:settings')


@crm_access_required
@require_POST
def sales_channel_restore_view(request, channel_id):
    sales_channel = _get_organization_object_or_404(request, SalesChannel, channel_id)
    sales_channel.is_active = True
    sales_channel.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Канал продаж восстановлен из архива.')
    return redirect('crm:settings')


@crm_access_required
def products_view(request):
    organization = request.crm_organization
    sku_query = (request.GET.get('sku') or '').strip()
    name_query = (request.GET.get('name') or '').strip()
    stock_from_raw = (request.GET.get('stock_from') or '').strip()
    stock_to_raw = (request.GET.get('stock_to') or '').strip()
    selected_sort = (request.GET.get('sort') or 'name_asc').strip()
    if selected_sort not in PRODUCT_SORT_OPTIONS:
        selected_sort = 'name_asc'

    try:
        stock_from = int(stock_from_raw) if stock_from_raw != '' else None
        if stock_from is not None and stock_from < 0:
            stock_from = None
    except ValueError:
        stock_from = None

    try:
        stock_to = int(stock_to_raw) if stock_to_raw != '' else None
        if stock_to is not None and stock_to < 0:
            stock_to = None
    except ValueError:
        stock_to = None

    lot_cost_expression = ExpressionWrapper(
        F('stock_lots__remaining_quantity') * (F('stock_lots__unit_cost') + F('stock_lots__extra_cost_per_unit')),
        output_field=DecimalField(max_digits=16, decimal_places=2),
    )
    products = (
        Product.objects.filter(organization=organization)
        .annotate(
            total_stock_quantity=Coalesce(
                Sum(
                    'stock_lots__remaining_quantity',
                    filter=Q(stock_lots__remaining_quantity__gt=0),
                    output_field=IntegerField(),
                ),
                Value(0),
                output_field=IntegerField(),
            ),
            stock_lot_count=Count(
                'stock_lots',
                filter=Q(stock_lots__remaining_quantity__gt=0),
                distinct=True,
            ),
            total_stock_cost=Coalesce(
                Sum(
                    lot_cost_expression,
                    filter=Q(stock_lots__remaining_quantity__gt=0),
                    output_field=DecimalField(max_digits=16, decimal_places=2),
                ),
                Value(Decimal('0.00')),
                output_field=DecimalField(max_digits=16, decimal_places=2),
            ),
        )
        .annotate(
            calculated_unit_cost=Case(
                When(
                    total_stock_quantity__gt=0,
                    then=ExpressionWrapper(
                        F('total_stock_cost') / F('total_stock_quantity'),
                        output_field=DecimalField(max_digits=16, decimal_places=2),
                    ),
                ),
                default=Value(None),
                output_field=DecimalField(max_digits=16, decimal_places=2),
            ),
        )
        .prefetch_related(
            Prefetch(
                'channel_offers',
                queryset=ProductChannelOffer.objects.select_related('channel').filter(
                    is_enabled=True,
                    channel__is_active=True,
                    channel__is_listing_channel=True,
                ),
                to_attr='enabled_listing_offers',
            )
        )
    )

    if sku_query:
        products = products.filter(sku__icontains=sku_query)
    if name_query:
        products = products.filter(name__icontains=name_query)
    if stock_from is not None:
        products = products.filter(total_stock_quantity__gte=stock_from)
    if stock_to is not None:
        products = products.filter(total_stock_quantity__lte=stock_to)

    products = products.order_by(*PRODUCT_SORT_OPTIONS[selected_sort]['ordering'])

    paginator = Paginator(products, PRODUCTS_PER_PAGE)
    page_obj = paginator.get_page(request.GET.get('page'))
    products = list(page_obj.object_list)
    for product in products:
        if product.total_stock_quantity:
            product.calculated_unit_cost = Decimal(product.calculated_unit_cost).quantize(Decimal('0.01'))
        else:
            product.calculated_unit_cost = None
    page_obj.object_list = products

    query_params = request.GET.copy()
    query_params.pop('page', None)

    return render(
        request,
        'crm/products.html',
        {
            'organization': organization,
            'products': products,
            'page_obj': page_obj,
            'pagination_query': query_params.urlencode(),
            'product_import_summary': _pop_product_import_summary(request),
            'filters': {
                'sku': sku_query,
                'name': name_query,
                'stock_from': stock_from_raw,
                'stock_to': stock_to_raw,
                'sort': selected_sort,
            },
            'sort_options': PRODUCT_SORT_OPTIONS,
        },
    )


@crm_access_required
def product_template_download_view(request):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Товары'
    worksheet.append(PRODUCT_IMPORT_TEMPLATE_HEADERS)
    worksheet.append(['SKU-001', 'Пример товара', 'Короткий комментарий', 'Вставьте фото в эту строку'])
    worksheet.column_dimensions['A'].width = 20
    worksheet.column_dimensions['B'].width = 32
    worksheet.column_dimensions['C'].width = 28
    worksheet.column_dimensions['D'].width = 20

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="crm-products-template.xlsx"'
    workbook.save(response)
    return response


@crm_access_required
def product_bulk_import_view(request):
    if request.method == 'POST':
        form = ProductBulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                rows = _parse_product_import_xlsx(
                    xlsx_file=form.cleaned_data['xlsx_file'],
                )
            except ValueError as error:
                form.add_error('xlsx_file', str(error))
            else:
                with transaction.atomic():
                    created_count, skipped_rows = _create_products_from_rows(
                        organization=request.crm_organization,
                        rows=rows,
                    )
                messages.success(request, f'Импорт завершен: загружено {created_count}.')
                if skipped_rows:
                    request.session[PRODUCT_IMPORT_SUMMARY_SESSION_KEY] = _build_product_import_summary(skipped_rows)
                else:
                    request.session.pop(PRODUCT_IMPORT_SUMMARY_SESSION_KEY, None)
                return redirect('crm:products')
    else:
        form = ProductBulkUploadForm()

    return render(
        request,
        'crm/product_import.html',
        {
            'organization': request.crm_organization,
            'form': form,
            'page_title': 'Массовая загрузка товаров',
        },
    )


@crm_access_required
def product_create_view(request):
    listing_channels = list(_get_listing_channels(request.crm_organization))

    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES, organization=request.crm_organization)
        offer_form = ProductChannelOfferForm(request.POST, channels=listing_channels)
        if form.is_valid() and offer_form.is_valid():
            with transaction.atomic():
                product = form.save(commit=False)
                product.organization = request.crm_organization
                product.save()
                _save_product_channel_offers(product, offer_form)
            messages.success(request, 'Товар создан.')
            return redirect('crm:products')
    else:
        form = ProductForm(initial={'is_active': True}, organization=request.crm_organization)
        offer_form = ProductChannelOfferForm(channels=listing_channels)

    return _render_product_form(
        request,
        form,
        offer_form,
        page_title='Новый товар',
        submit_label='Создать товар',
        stock_lots=[],
    )


@crm_access_required
def product_edit_view(request, product_id):
    product = _get_organization_object_or_404(request, Product, product_id)
    listing_channels = list(_get_listing_channels(request.crm_organization))
    stock_lots = _get_product_stock_lots(product)
    existing_offers = list(
        ProductChannelOffer.objects.filter(
            product=product,
            channel__organization=request.crm_organization,
            channel__is_listing_channel=True,
        ).select_related('channel')
    )

    if request.method == 'POST':
        form = ProductForm(
            request.POST,
            request.FILES,
            instance=product,
            organization=request.crm_organization,
        )
        offer_form = ProductChannelOfferForm(
            request.POST,
            channels=listing_channels,
            offers=existing_offers,
        )
        if form.is_valid() and offer_form.is_valid():
            with transaction.atomic():
                form.save()
                _save_product_channel_offers(product, offer_form)
            messages.success(request, 'Товар обновлен.')
            return redirect('crm:products')
    else:
        form = ProductForm(instance=product, organization=request.crm_organization)
        offer_form = ProductChannelOfferForm(channels=listing_channels, offers=existing_offers)

    return _render_product_form(
        request,
        form,
        offer_form,
        page_title=f'Редактирование товара: {product.name}',
        submit_label='Сохранить',
        stock_lots=stock_lots,
    )


@crm_access_required
@require_POST
def product_archive_view(request, product_id):
    product = _get_organization_object_or_404(request, Product, product_id)
    product.is_active = False
    product.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Товар отправлен в архив.')
    return redirect('crm:products')


@crm_access_required
def orders_view(request):
    organization = request.crm_organization
    query = (request.GET.get('q') or '').strip()
    status_filter = (request.GET.get('status') or '').strip()
    channel_filter = (request.GET.get('sales_channel') or '').strip()
    date_from = (request.GET.get('date_from') or '').strip()
    date_to = (request.GET.get('date_to') or '').strip()

    line_total_expression = ExpressionWrapper(
        F('items__quantity') * Coalesce(
            F('items__sale_price'),
            Value(Decimal('0.00')),
            output_field=DecimalField(max_digits=12, decimal_places=2),
        ),
        output_field=DecimalField(max_digits=14, decimal_places=2),
    )
    orders = (
        Order.objects.filter(organization=organization)
        .select_related('sales_channel', 'delivery_method')
        .annotate(
            line_count=Coalesce(Count('items', distinct=True), Value(0), output_field=IntegerField()),
            total_quantity=Coalesce(Sum('items__quantity'), Value(0), output_field=IntegerField()),
            total_amount=Coalesce(
                Sum(line_total_expression, output_field=DecimalField(max_digits=14, decimal_places=2)),
                Value(Decimal('0.00')),
                output_field=DecimalField(max_digits=14, decimal_places=2),
            ),
        )
        .order_by('-order_date', '-id')
    )

    if query:
        orders = orders.filter(
            Q(internal_order_number__icontains=query)
            | Q(external_order_number__icontains=query)
            | Q(customer_name__icontains=query)
            | Q(customer_phone__icontains=query)
            | Q(customer_contact__icontains=query)
        )
    if status_filter:
        orders = orders.filter(status=status_filter)
    if channel_filter:
        orders = orders.filter(sales_channel_id=channel_filter)
    if date_from:
        orders = orders.filter(order_date__gte=date_from)
    if date_to:
        orders = orders.filter(order_date__lte=date_to)

    paginator = Paginator(orders, ORDERS_PER_PAGE)
    page_obj = paginator.get_page(request.GET.get('page'))
    orders = list(page_obj.object_list)
    for order in orders:
        _apply_order_summary(order)
    page_obj.object_list = orders

    query_params = request.GET.copy()
    query_params.pop('page', None)

    return render(
        request,
        'crm/orders.html',
        {
            'organization': organization,
            'orders': orders,
            'page_obj': page_obj,
            'pagination_query': query_params.urlencode(),
            'sales_channels': _get_order_sales_channels_queryset(organization),
            'status_choices': Order.Status.choices,
            'order_status_badges': ORDER_STATUS_BADGES,
            'filters': {
                'q': query,
                'status': status_filter,
                'sales_channel': channel_filter,
                'date_from': date_from,
                'date_to': date_to,
            },
        },
    )


@crm_access_required
def order_create_view(request):
    order = Order(organization=request.crm_organization, created_by=request.user)

    if request.method == 'POST':
        form = OrderForm(
            request.POST,
            instance=order,
            organization=request.crm_organization,
            allow_status_edit=False,
        )
        item_formset = _build_order_item_formset(request, order=order, data=request.POST)
        if form.is_valid() and item_formset.is_valid():
            with transaction.atomic():
                order = form.save(commit=False)
                order.organization = request.crm_organization
                order.created_by = request.user
                order.status = Order.Status.NEW
                order.save()
                result = replace_order_items_and_reservations(order, item_formset.get_cleaned_rows())

            if result['shortage_item_ids']:
                messages.warning(request, 'Не удалось полностью зарезервировать остатки по всем строкам заказа.')
            if result['missing_price_product_ids']:
                messages.warning(request, 'Для части товаров цена по выбранному каналу не найдена и осталась пустой.')
            messages.success(request, 'Заказ создан.')
            return redirect('crm:order_detail', order_id=order.id)
    else:
        form = OrderForm(instance=order, organization=request.crm_organization, allow_status_edit=False)
        item_formset = _build_order_item_formset(request, order=order)

    return _render_order_form(
        request,
        form,
        item_formset,
        page_title='Новый заказ',
        submit_label='Создать заказ',
        order=None,
    )


@crm_access_required
def order_detail_view(request, order_id):
    order = _get_organization_object_or_404(request, Order, order_id)
    order_summary = _apply_order_summary(order)
    return render(
        request,
        'crm/order_detail.html',
        {
            'organization': request.crm_organization,
            'order': order,
            'order_summary': order_summary,
            'order_status_badges': ORDER_STATUS_BADGES,
        },
    )


@crm_access_required
def order_edit_view(request, order_id):
    order = _get_organization_object_or_404(request, Order, order_id)

    if request.method == 'POST':
        form = OrderForm(
            request.POST,
            instance=order,
            organization=request.crm_organization,
            allow_status_edit=True,
        )
        item_formset = _build_order_item_formset(request, order=order, data=request.POST)
        if form.is_valid() and item_formset.is_valid():
            with transaction.atomic():
                order = form.save()
                result = replace_order_items_and_reservations(order, item_formset.get_cleaned_rows())

            if order.status == Order.Status.CANCELLED:
                messages.warning(request, 'Заказ сохранен в статусе "Отменен". Активные резервы не создаются.')
            elif result['shortage_item_ids']:
                messages.warning(request, 'После пересчета резервов не хватило остатков по части товаров.')
            if result['missing_price_product_ids']:
                messages.warning(request, 'Для части товаров цена по выбранному каналу не найдена и осталась пустой.')
            messages.success(request, 'Заказ обновлен.')
            return redirect('crm:order_detail', order_id=order.id)
    else:
        form = OrderForm(instance=order, organization=request.crm_organization, allow_status_edit=True)
        item_formset = _build_order_item_formset(
            request,
            order=order,
            initial=_get_order_item_initial_rows(order),
        )

    return _render_order_form(
        request,
        form,
        item_formset,
        page_title=f'Редактирование заказа: {order.internal_order_number}',
        submit_label='Сохранить',
        order=order,
    )


@crm_access_required
@require_POST
def order_cancel_view(request, order_id):
    order = _get_organization_object_or_404(request, Order, order_id)
    with transaction.atomic():
        order.status = Order.Status.CANCELLED
        order.save(update_fields=['status', 'updated_at'])
        release_order_reservations(order)
    messages.success(request, 'Заказ отменен, активные резервы освобождены.')
    return redirect('crm:orders')


@crm_access_required
def purchases_view(request):
    organization = request.crm_organization
    line_total_expression = (
        F('items__quantity') * F('items__unit_cost')
    ) + F('items__extra_cost_total')
    purchases = (
        Purchase.objects.filter(organization=organization)
        .annotate(
            line_count=Coalesce(Count('items'), Value(0), output_field=IntegerField()),
            total_quantity=Coalesce(Sum('items__quantity'), Value(0), output_field=IntegerField()),
            total_amount=Coalesce(
                Sum(line_total_expression, output_field=DecimalField(max_digits=14, decimal_places=2)),
                Value(Decimal('0.00')),
                output_field=DecimalField(max_digits=14, decimal_places=2),
            ),
        )
        .order_by('-purchase_date', '-id')
    )

    return render(
        request,
        'crm/purchases.html',
        {
            'organization': organization,
            'purchases': purchases,
        },
    )


@crm_access_required
def purchase_template_download_view(request):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="crm-purchases-template.csv"'
    response.write('\ufeff')
    writer = csv.writer(response)
    writer.writerow(PURCHASE_TEMPLATE_HEADERS)
    writer.writerow(['SKU-001', '10', '7000', '300', 'Первая партия'])
    return response


@crm_access_required
def purchase_create_view(request):
    purchase = Purchase(organization=request.crm_organization, created_by=request.user)

    if request.method == 'POST':
        if 'load_csv' in request.POST:
            form = PurchaseForm(request.POST, instance=purchase, organization=request.crm_organization)
            upload_form = PurchaseCSVUploadForm(request.POST, request.FILES)
            csv_preview = None
            if upload_form.is_valid() and upload_form.cleaned_data.get('file'):
                try:
                    warehouse = _resolve_purchase_import_warehouse(request.POST.get('warehouse'), request.crm_organization)
                    parsed_rows = _parse_purchase_csv(upload_form.cleaned_data['file'], request.crm_organization)
                    csv_preview = _build_purchase_csv_preview(parsed_rows, request.crm_organization, warehouse)
                    messages.success(request, 'CSV проверен. Строки подгружены в форму, но закупка еще не создана.')
                    item_formset = _build_purchase_item_formset(request, purchase=purchase, initial=parsed_rows)
                except ValueError as error:
                    error_text = str(error)
                    if 'склад' in error_text.lower():
                        form.add_error('warehouse', error_text)
                    else:
                        upload_form.add_error('file', error_text)
                    item_formset = _build_purchase_item_formset(request, purchase=purchase)
            else:
                if not upload_form.cleaned_data.get('file'):
                    upload_form.add_error('file', 'Выберите CSV-файл по шаблону.')
                item_formset = _build_purchase_item_formset(request, purchase=purchase)

            return _render_purchase_form(
                request,
                form,
                item_formset,
                upload_form,
                page_title='Новая закупка',
                submit_label='Создать закупку',
                csv_preview=csv_preview,
            )

        form = PurchaseForm(request.POST, instance=purchase, organization=request.crm_organization)
        item_formset = _build_purchase_item_formset(request, purchase=purchase, data=request.POST)
        upload_form = PurchaseCSVUploadForm()
        if form.is_valid() and item_formset.is_valid():
            with transaction.atomic():
                purchase = form.save(commit=False)
                purchase.organization = request.crm_organization
                purchase.created_by = request.user
                purchase.save()
                replace_purchase_items_and_stock(
                    purchase,
                    item_formset.get_cleaned_rows(),
                    created_by=request.user,
                )
            messages.success(request, 'Закупка создана.')
            return redirect('crm:purchases')
    else:
        form = PurchaseForm(instance=purchase, organization=request.crm_organization)
        item_formset = _build_purchase_item_formset(request, purchase=purchase)
        upload_form = PurchaseCSVUploadForm()

    return _render_purchase_form(
        request,
        form,
        item_formset,
        upload_form,
        page_title='Новая закупка',
        submit_label='Создать закупку',
    )


@crm_access_required
def purchase_detail_view(request, purchase_id):
    purchase = _get_organization_object_or_404(request, Purchase, purchase_id)
    detail_summary = _get_purchase_detail_summary(purchase)
    return render(
        request,
        'crm/purchase_detail.html',
        {
            'organization': request.crm_organization,
            'purchase': purchase,
            'purchase_items': detail_summary['items'],
            'purchase_line_count': detail_summary['line_count'],
            'purchase_total_quantity': detail_summary['total_quantity'],
            'purchase_total_amount': detail_summary['total_amount'],
        },
    )


@crm_access_required
def purchase_edit_view(request, purchase_id):
    purchase = _get_organization_object_or_404(request, Purchase, purchase_id)

    if request.method == 'POST':
        if 'load_csv' in request.POST:
            form = PurchaseForm(request.POST, instance=purchase, organization=request.crm_organization)
            upload_form = PurchaseCSVUploadForm(request.POST, request.FILES)
            csv_preview = None
            if upload_form.is_valid() and upload_form.cleaned_data.get('file'):
                try:
                    warehouse = _resolve_purchase_import_warehouse(
                        request.POST.get('warehouse') or str(purchase.warehouse_id or ''),
                        request.crm_organization,
                    )
                    parsed_rows = _parse_purchase_csv(upload_form.cleaned_data['file'], request.crm_organization)
                    csv_preview = _build_purchase_csv_preview(parsed_rows, request.crm_organization, warehouse)
                    messages.success(request, 'CSV проверен. Строки подгружены в форму, но закупка еще не сохранена.')
                    item_formset = _build_purchase_item_formset(request, purchase=purchase, initial=parsed_rows)
                except ValueError as error:
                    error_text = str(error)
                    if 'склад' in error_text.lower():
                        form.add_error('warehouse', error_text)
                    else:
                        upload_form.add_error('file', error_text)
                    item_formset = _build_purchase_item_formset(
                        request,
                        purchase=purchase,
                        initial=_get_purchase_item_initial_rows(purchase),
                    )
            else:
                if not upload_form.cleaned_data.get('file'):
                    upload_form.add_error('file', 'Выберите CSV-файл по шаблону.')
                item_formset = _build_purchase_item_formset(
                    request,
                    purchase=purchase,
                    initial=_get_purchase_item_initial_rows(purchase),
                )

            return _render_purchase_form(
                request,
                form,
                item_formset,
                upload_form,
                page_title=f'Редактирование закупки от {purchase.purchase_date:%d.%m.%Y}',
                submit_label='Сохранить',
                purchase=purchase,
                csv_preview=csv_preview,
            )

        form = PurchaseForm(request.POST, instance=purchase, organization=request.crm_organization)
        item_formset = _build_purchase_item_formset(request, purchase=purchase, data=request.POST)
        upload_form = PurchaseCSVUploadForm()
        if form.is_valid() and item_formset.is_valid():
            with transaction.atomic():
                purchase = form.save()
                replace_purchase_items_and_stock(
                    purchase,
                    item_formset.get_cleaned_rows(),
                    created_by=request.user,
                )
            messages.success(request, 'Закупка обновлена.')
            return redirect('crm:purchases')
    else:
        form = PurchaseForm(instance=purchase, organization=request.crm_organization)
        item_formset = _build_purchase_item_formset(
            request,
            purchase=purchase,
            initial=_get_purchase_item_initial_rows(purchase),
        )
        upload_form = PurchaseCSVUploadForm()

    return _render_purchase_form(
        request,
        form,
        item_formset,
        upload_form,
        page_title=f'Редактирование закупки от {purchase.purchase_date:%d.%m.%Y}',
        submit_label='Сохранить',
        purchase=purchase,
    )


@crm_access_required
@require_POST
def purchase_archive_view(request, purchase_id):
    purchase = _get_organization_object_or_404(request, Purchase, purchase_id)
    with transaction.atomic():
        purchase.status = Purchase.Status.CANCELLED
        purchase.save(update_fields=['status', 'updated_at'])
        replace_purchase_items_and_stock(
            purchase,
            [
                {
                    'product': item.product,
                    'warehouse': item.warehouse,
                    'quantity': item.quantity,
                    'unit_cost': item.unit_cost,
                    'extra_cost_total': item.extra_cost_total,
                    'comment': item.comment,
                }
                for item in purchase.items.select_related('product', 'warehouse')
            ],
            created_by=request.user,
        )
    messages.success(request, 'Закупка отменена.')
    return redirect('crm:purchases')


@crm_access_required
def warehouse_create_view(request):
    if request.method == 'POST':
        form = WarehouseForm(request.POST)
        if form.is_valid():
            warehouse = form.save(commit=False)
            warehouse.organization = request.crm_organization
            warehouse.save()
            messages.success(request, 'Склад создан.')
            return redirect('crm:settings')
    else:
        form = WarehouseForm()

    return _render_directory_form(
        request,
        form,
        page_title='Новый склад',
        submit_label='Создать склад',
        cancel_url='crm:settings',
    )


@crm_access_required
def warehouse_edit_view(request, warehouse_id):
    warehouse = _get_organization_object_or_404(request, Warehouse, warehouse_id)

    if request.method == 'POST':
        form = WarehouseForm(request.POST, instance=warehouse)
        if form.is_valid():
            form.save()
            messages.success(request, 'Склад обновлен.')
            return redirect('crm:settings')
    else:
        form = WarehouseForm(instance=warehouse)

    return _render_directory_form(
        request,
        form,
        page_title=f'Редактирование склада: {warehouse.name}',
        submit_label='Сохранить',
        cancel_url='crm:settings',
    )


@crm_access_required
@require_POST
def warehouse_archive_view(request, warehouse_id):
    warehouse = _get_organization_object_or_404(request, Warehouse, warehouse_id)
    warehouse.is_active = False
    warehouse.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Склад отправлен в архив.')
    return redirect('crm:settings')


@crm_access_required
@require_POST
def warehouse_restore_view(request, warehouse_id):
    warehouse = _get_organization_object_or_404(request, Warehouse, warehouse_id)
    warehouse.is_active = True
    warehouse.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Склад восстановлен из архива.')
    return redirect('crm:settings')


@crm_access_required
def delivery_method_create_view(request):
    if request.method == 'POST':
        form = DeliveryMethodForm(request.POST, organization=request.crm_organization)
        if form.is_valid():
            delivery_method = form.save(commit=False)
            delivery_method.organization = request.crm_organization
            delivery_method.save()
            messages.success(request, 'Способ доставки создан.')
            return redirect('crm:settings')
    else:
        form = DeliveryMethodForm(organization=request.crm_organization)

    return _render_directory_form(
        request,
        form,
        page_title='Новый способ доставки',
        submit_label='Создать способ доставки',
        cancel_url='crm:settings',
    )


@crm_access_required
def delivery_method_edit_view(request, delivery_method_id):
    delivery_method = _get_organization_object_or_404(request, DeliveryMethod, delivery_method_id)

    if request.method == 'POST':
        form = DeliveryMethodForm(request.POST, instance=delivery_method, organization=request.crm_organization)
        if form.is_valid():
            form.save()
            messages.success(request, 'Способ доставки обновлен.')
            return redirect('crm:settings')
    else:
        form = DeliveryMethodForm(instance=delivery_method, organization=request.crm_organization)

    return _render_directory_form(
        request,
        form,
        page_title=f'Редактирование способа доставки: {delivery_method.name}',
        submit_label='Сохранить',
        cancel_url='crm:settings',
    )


@crm_access_required
@require_POST
def delivery_method_archive_view(request, delivery_method_id):
    delivery_method = _get_organization_object_or_404(request, DeliveryMethod, delivery_method_id)
    delivery_method.is_active = False
    delivery_method.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Способ доставки отправлен в архив.')
    return redirect('crm:settings')


@crm_access_required
@require_POST
def delivery_method_restore_view(request, delivery_method_id):
    delivery_method = _get_organization_object_or_404(request, DeliveryMethod, delivery_method_id)
    delivery_method.is_active = True
    delivery_method.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Способ доставки восстановлен из архива.')
    return redirect('crm:settings')


@crm_access_required
def counterparty_create_view(request):
    if request.method == 'POST':
        form = CounterpartyForm(request.POST)
        if form.is_valid():
            counterparty = form.save(commit=False)
            counterparty.organization = request.crm_organization
            counterparty.save()
            messages.success(request, 'Контрагент создан.')
            return redirect('crm:settings')
    else:
        form = CounterpartyForm()

    return _render_directory_form(
        request,
        form,
        page_title='Новый контрагент',
        submit_label='Создать контрагента',
        cancel_url='crm:settings',
    )


@crm_access_required
def counterparty_edit_view(request, counterparty_id):
    counterparty = _get_organization_object_or_404(request, Counterparty, counterparty_id)

    if request.method == 'POST':
        form = CounterpartyForm(request.POST, instance=counterparty)
        if form.is_valid():
            form.save()
            messages.success(request, 'Контрагент обновлен.')
            return redirect('crm:settings')
    else:
        form = CounterpartyForm(instance=counterparty)

    return _render_directory_form(
        request,
        form,
        page_title=f'Редактирование контрагента: {counterparty.name}',
        submit_label='Сохранить',
        cancel_url='crm:settings',
    )


@crm_access_required
@require_POST
def counterparty_archive_view(request, counterparty_id):
    counterparty = _get_organization_object_or_404(request, Counterparty, counterparty_id)
    counterparty.is_active = False
    counterparty.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Контрагент отправлен в архив.')
    return redirect('crm:settings')


@crm_access_required
@require_POST
def counterparty_restore_view(request, counterparty_id):
    counterparty = _get_organization_object_or_404(request, Counterparty, counterparty_id)
    counterparty.is_active = True
    counterparty.save(update_fields=['is_active', 'updated_at'])
    messages.success(request, 'Контрагент восстановлен из архива.')
    return redirect('crm:settings')


@login_required
def onboarding_view(request):
    current_organization = get_current_organization(request.user)
    if current_organization is not None:
        return redirect('crm:dashboard')

    if request.method == 'POST':
        form = OrganizationOnboardingForm(request.POST)
        if form.is_valid():
            with transaction.atomic():
                organization = form.save(commit=False)
                organization.owner = request.user
                organization.save()
                Membership.objects.create(
                    organization=organization,
                    user=request.user,
                    role=Membership.Role.OWNER,
                )
                initialize_organization_defaults(organization)
            return redirect('crm:dashboard')
    else:
        form = OrganizationOnboardingForm()

    return render(request, 'crm/onboarding.html', {'form': form})
