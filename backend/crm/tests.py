import io
import os
import tempfile

from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from django.db.models import Sum
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image as PILImage

from .access import get_current_organization, user_has_crm_access
from .models import (
    Counterparty,
    DeliveryMethod,
    Membership,
    Order,
    OrderItem,
    Organization,
    Product,
    ProductChannelOffer,
    Purchase,
    PurchaseItem,
    SalesChannel,
    StockLot,
    StockMovement,
    StockReservation,
    Warehouse,
)
from .services import (
    DEFAULT_DELIVERY_METHODS,
    DEFAULT_SALES_CHANNELS,
    DEFAULT_WAREHOUSES,
    PRODUCT_IMPORT_TEMPLATE_HEADERS,
    PURCHASE_TEMPLATE_HEADERS,
    release_order_reservations,
    replace_order_items_and_reservations,
    replace_purchase_items_and_stock,
)


TEST_GIF_BYTES = (
    b'GIF89a\x01\x00\x01\x00\x80\x00\x00'
    b'\x00\x00\x00\xff\xff\xff!\xf9\x04'
    b'\x01\x00\x00\x00\x00,\x00\x00\x00'
    b'\x00\x01\x00\x01\x00\x00\x02\x02D'
    b'\x01\x00;'
)


def build_product_import_workbook(headers, rows, image_cells=None):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Товары'
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    image_cells = image_cells or []
    temporary_image_paths = []
    for cell_address in image_cells:
        with tempfile.NamedTemporaryFile(suffix='.png', delete=False) as temp_image:
            temporary_image_paths.append(temp_image.name)
            image = PILImage.new('RGB', (16, 16), color=(220, 30, 30))
            image.save(temp_image.name)
        worksheet.add_image(ExcelImage(temp_image.name), cell_address)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    for temp_image_path in temporary_image_paths:
        os.unlink(temp_image_path)

    return output.getvalue()


class CRMOnboardingFlowTest(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='crm-owner', password='password123')

    def test_dashboard_redirects_to_onboarding_without_active_membership(self):
        self.client.login(username='crm-owner', password='password123')

        response = self.client.get(reverse('crm:dashboard'))

        self.assertRedirects(response, reverse('crm:onboarding'))

    def test_onboarding_creates_organization_and_owner_membership(self):
        self.client.login(username='crm-owner', password='password123')

        response = self.client.post(
            reverse('crm:onboarding'),
            {
                'name': 'LEGO бизнес',
                'base_currency': 'RUB',
            },
        )

        self.assertRedirects(response, reverse('crm:dashboard'))
        organization = Organization.objects.get()
        membership = Membership.objects.get()

        self.assertEqual(organization.owner, self.user)
        self.assertEqual(organization.name, 'LEGO бизнес')
        self.assertEqual(organization.base_currency, 'RUB')
        self.assertEqual(membership.organization, organization)
        self.assertEqual(membership.user, self.user)
        self.assertEqual(membership.role, Membership.Role.OWNER)
        self.assertCountEqual(
            SalesChannel.objects.filter(organization=organization).values_list('name', flat=True),
            [item['name'] for item in DEFAULT_SALES_CHANNELS],
        )
        self.assertCountEqual(
            Warehouse.objects.filter(organization=organization).values_list('name', flat=True),
            DEFAULT_WAREHOUSES,
        )
        self.assertCountEqual(
            DeliveryMethod.objects.filter(organization=organization).values_list('name', flat=True),
            [item['name'] for item in DEFAULT_DELIVERY_METHODS],
        )
        self.assertTrue(
            SalesChannel.objects.get(organization=organization, name='Ozon').is_listing_channel
        )
        self.assertTrue(
            SalesChannel.objects.get(organization=organization, name='Яндекс Маркет').is_listing_channel
        )
        self.assertTrue(
            SalesChannel.objects.get(organization=organization, name='Avito').is_listing_channel
        )
        self.assertFalse(
            SalesChannel.objects.get(organization=organization, name='Самовывоз').is_listing_channel
        )

        dashboard_response = self.client.get(reverse('crm:dashboard'))
        self.assertContains(dashboard_response, 'LEGO бизнес')
        self.assertContains(dashboard_response, 'Справочники')
        self.assertNotContains(dashboard_response, 'Аналитика')


class CRMAccessHelpersTest(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='crm-owner', password='password123')
        self.outsider = User.objects.create_user(username='crm-outsider', password='password123')
        self.organization = Organization.objects.create(
            owner=self.owner,
            name='Brick Trade',
            base_currency='RUB',
        )
        Membership.objects.create(
            organization=self.organization,
            user=self.owner,
            role=Membership.Role.OWNER,
        )

    def test_get_current_organization_returns_first_active_organization(self):
        organization = get_current_organization(self.owner)

        self.assertEqual(organization, self.organization)

    def test_user_has_crm_access_denies_non_member(self):
        self.assertFalse(user_has_crm_access(self.outsider, self.organization))

    def test_dashboard_does_not_expose_other_users_organization(self):
        self.client.login(username='crm-outsider', password='password123')

        response = self.client.get(reverse('crm:dashboard'))

        self.assertRedirects(response, reverse('crm:onboarding'))

    def test_membership_unique_constraint_blocks_duplicate_user_in_organization(self):
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Membership.objects.create(
                    organization=self.organization,
                    user=self.owner,
                    role=Membership.Role.ADMIN,
                )

    def test_finance_pages_use_finance_sidebar(self):
        self.client.login(username='crm-owner', password='password123')

        response = self.client.get(reverse('profile'))

        self.assertContains(response, 'Аналитика')
        self.assertContains(response, 'Финансы')
        self.assertContains(response, 'CRM')
        self.assertNotContains(response, 'Инвентаризация')


class CRMSettingsDirectoriesTest(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='crm-settings-owner', password='password123')
        self.other_owner = User.objects.create_user(username='crm-settings-other', password='password123')

        self.organization = Organization.objects.create(
            owner=self.owner,
            name='Brick Trade',
            base_currency='RUB',
        )
        self.other_organization = Organization.objects.create(
            owner=self.other_owner,
            name='Other Org',
            base_currency='RUB',
        )

        Membership.objects.create(
            organization=self.organization,
            user=self.owner,
            role=Membership.Role.OWNER,
        )
        Membership.objects.create(
            organization=self.other_organization,
            user=self.other_owner,
            role=Membership.Role.OWNER,
        )

        self.sales_channel = SalesChannel.objects.create(
            organization=self.organization,
            name='Ozon',
        )
        self.archived_sales_channel = SalesChannel.objects.create(
            organization=self.organization,
            name='Старый канал',
            is_active=False,
        )
        self.other_sales_channel = SalesChannel.objects.create(
            organization=self.other_organization,
            name='Чужой канал',
        )

        self.warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Основной склад',
        )
        self.archived_warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Старый склад',
            is_active=False,
        )
        self.other_warehouse = Warehouse.objects.create(
            organization=self.other_organization,
            name='Чужой склад',
        )
        self.counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Поставщик 1',
        )
        self.archived_counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Старый контрагент',
            is_active=False,
        )
        self.other_counterparty = Counterparty.objects.create(
            organization=self.other_organization,
            name='Чужой контрагент',
        )

    def test_settings_page_shows_only_current_organization_records(self):
        self.client.login(username='crm-settings-owner', password='password123')

        response = self.client.get(reverse('crm:settings'))

        self.assertContains(response, 'Ozon')
        self.assertContains(response, 'Старый канал')
        self.assertContains(response, 'Основной склад')
        self.assertContains(response, 'Старый склад')
        self.assertContains(response, 'Поставщик 1')
        self.assertContains(response, 'Старый контрагент')
        self.assertNotContains(response, 'Чужой канал')
        self.assertNotContains(response, 'Чужой склад')
        self.assertNotContains(response, 'Чужой контрагент')

    def test_user_can_create_edit_archive_and_restore_sales_channel_in_own_organization(self):
        self.client.login(username='crm-settings-owner', password='password123')

        create_response = self.client.post(
            reverse('crm:sales_channel_create'),
            {'name': 'Wildberries', 'is_listing_channel': 'on'},
        )
        self.assertRedirects(create_response, reverse('crm:settings'))

        created_channel = SalesChannel.objects.get(name='Wildberries')
        self.assertEqual(created_channel.organization, self.organization)
        self.assertTrue(created_channel.is_active)
        self.assertTrue(created_channel.is_listing_channel)

        edit_response = self.client.post(
            reverse('crm:sales_channel_edit', args=[created_channel.id]),
            {'name': 'WB'},
        )
        self.assertRedirects(edit_response, reverse('crm:settings'))
        created_channel.refresh_from_db()
        self.assertEqual(created_channel.name, 'WB')
        self.assertFalse(created_channel.is_listing_channel)

        archive_response = self.client.post(
            reverse('crm:sales_channel_archive', args=[created_channel.id]),
        )
        self.assertRedirects(archive_response, reverse('crm:settings'))
        created_channel.refresh_from_db()
        self.assertFalse(created_channel.is_active)

        restore_response = self.client.post(
            reverse('crm:sales_channel_restore', args=[created_channel.id]),
        )
        self.assertRedirects(restore_response, reverse('crm:settings'))
        created_channel.refresh_from_db()
        self.assertTrue(created_channel.is_active)

    def test_user_can_create_edit_archive_and_restore_warehouse_in_own_organization(self):
        self.client.login(username='crm-settings-owner', password='password123')

        create_response = self.client.post(
            reverse('crm:warehouse_create'),
            {'name': 'Региональный склад'},
        )
        self.assertRedirects(create_response, reverse('crm:settings'))

        warehouse = Warehouse.objects.get(name='Региональный склад')
        self.assertEqual(warehouse.organization, self.organization)
        self.assertTrue(warehouse.is_active)

        edit_response = self.client.post(
            reverse('crm:warehouse_edit', args=[warehouse.id]),
            {'name': 'Склад Север'},
        )
        self.assertRedirects(edit_response, reverse('crm:settings'))
        warehouse.refresh_from_db()
        self.assertEqual(warehouse.name, 'Склад Север')

        archive_response = self.client.post(
            reverse('crm:warehouse_archive', args=[warehouse.id]),
        )
        self.assertRedirects(archive_response, reverse('crm:settings'))
        warehouse.refresh_from_db()
        self.assertFalse(warehouse.is_active)

        restore_response = self.client.post(
            reverse('crm:warehouse_restore', args=[warehouse.id]),
        )
        self.assertRedirects(restore_response, reverse('crm:settings'))
        warehouse.refresh_from_db()
        self.assertTrue(warehouse.is_active)

    def test_user_can_create_edit_archive_and_restore_counterparty_in_own_organization(self):
        self.client.login(username='crm-settings-owner', password='password123')

        create_response = self.client.post(
            reverse('crm:counterparty_create'),
            {'name': 'Поставщик LEGO', 'comment': 'Основной поставщик'},
        )
        self.assertRedirects(create_response, reverse('crm:settings'))

        counterparty = Counterparty.objects.get(name='Поставщик LEGO')
        self.assertEqual(counterparty.organization, self.organization)
        self.assertTrue(counterparty.is_active)
        self.assertEqual(counterparty.comment, 'Основной поставщик')

        edit_response = self.client.post(
            reverse('crm:counterparty_edit', args=[counterparty.id]),
            {'name': 'Поставщик LEGO 2', 'comment': 'Обновили'},
        )
        self.assertRedirects(edit_response, reverse('crm:settings'))
        counterparty.refresh_from_db()
        self.assertEqual(counterparty.name, 'Поставщик LEGO 2')
        self.assertEqual(counterparty.comment, 'Обновили')

        archive_response = self.client.post(
            reverse('crm:counterparty_archive', args=[counterparty.id]),
        )
        self.assertRedirects(archive_response, reverse('crm:settings'))
        counterparty.refresh_from_db()
        self.assertFalse(counterparty.is_active)

        restore_response = self.client.post(
            reverse('crm:counterparty_restore', args=[counterparty.id]),
        )
        self.assertRedirects(restore_response, reverse('crm:settings'))
        counterparty.refresh_from_db()
        self.assertTrue(counterparty.is_active)

    def test_user_cannot_edit_archive_or_restore_other_organization_records(self):
        self.client.login(username='crm-settings-owner', password='password123')

        edit_channel_response = self.client.post(
            reverse('crm:sales_channel_edit', args=[self.other_sales_channel.id]),
            {'name': 'Попытка'},
        )
        self.assertEqual(edit_channel_response.status_code, 404)

        archive_channel_response = self.client.post(
            reverse('crm:sales_channel_archive', args=[self.other_sales_channel.id]),
        )
        self.assertEqual(archive_channel_response.status_code, 404)

        restore_channel_response = self.client.post(
            reverse('crm:sales_channel_restore', args=[self.other_sales_channel.id]),
        )
        self.assertEqual(restore_channel_response.status_code, 404)

        edit_warehouse_response = self.client.post(
            reverse('crm:warehouse_edit', args=[self.other_warehouse.id]),
            {'name': 'Попытка'},
        )
        self.assertEqual(edit_warehouse_response.status_code, 404)

        archive_warehouse_response = self.client.post(
            reverse('crm:warehouse_archive', args=[self.other_warehouse.id]),
        )
        self.assertEqual(archive_warehouse_response.status_code, 404)

        restore_warehouse_response = self.client.post(
            reverse('crm:warehouse_restore', args=[self.other_warehouse.id]),
        )
        self.assertEqual(restore_warehouse_response.status_code, 404)

        edit_counterparty_response = self.client.post(
            reverse('crm:counterparty_edit', args=[self.other_counterparty.id]),
            {'name': 'Попытка', 'comment': ''},
        )
        self.assertEqual(edit_counterparty_response.status_code, 404)

        archive_counterparty_response = self.client.post(
            reverse('crm:counterparty_archive', args=[self.other_counterparty.id]),
        )
        self.assertEqual(archive_counterparty_response.status_code, 404)

        restore_counterparty_response = self.client.post(
            reverse('crm:counterparty_restore', args=[self.other_counterparty.id]),
        )
        self.assertEqual(restore_counterparty_response.status_code, 404)


class CRMProductsTest(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='crm-products-owner', password='password123')
        self.other_owner = User.objects.create_user(username='crm-products-other', password='password123')

        self.organization = Organization.objects.create(
            owner=self.owner,
            name='Brick Trade',
            base_currency='RUB',
        )
        self.other_organization = Organization.objects.create(
            owner=self.other_owner,
            name='Other Org',
            base_currency='RUB',
        )

        Membership.objects.create(
            organization=self.organization,
            user=self.owner,
            role=Membership.Role.OWNER,
        )
        Membership.objects.create(
            organization=self.other_organization,
            user=self.other_owner,
            role=Membership.Role.OWNER,
        )

        self.ozon = SalesChannel.objects.create(
            organization=self.organization,
            name='Ozon',
            is_listing_channel=True,
        )
        self.avito = SalesChannel.objects.create(
            organization=self.organization,
            name='Avito',
            is_listing_channel=True,
        )
        self.pickup = SalesChannel.objects.create(
            organization=self.organization,
            name='Самовывоз',
            is_listing_channel=False,
        )
        self.archived_listing_channel = SalesChannel.objects.create(
            organization=self.organization,
            name='Старый маркетплейс',
            is_listing_channel=True,
            is_active=False,
        )
        self.other_listing_channel = SalesChannel.objects.create(
            organization=self.other_organization,
            name='Чужой канал',
            is_listing_channel=True,
        )

        self.product = Product.objects.create(
            organization=self.organization,
            sku='SKU-001',
            name='Базовый товар',
            comment='Тестовый комментарий',
        )
        ProductChannelOffer.objects.create(
            product=self.product,
            channel=self.ozon,
            is_enabled=True,
            price='9990.00',
            comment='Основной канал',
        )
        ProductChannelOffer.objects.create(
            product=self.product,
            channel=self.avito,
            is_enabled=False,
            price=None,
            comment='Пока выключен',
        )

        self.other_product = Product.objects.create(
            organization=self.other_organization,
            sku='SKU-999',
            name='Чужой товар',
        )
        ProductChannelOffer.objects.create(
            product=self.other_product,
            channel=self.other_listing_channel,
            is_enabled=True,
            price='555.00',
        )
        self.warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Основной склад',
        )
        self.counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Поставщик для товаров',
        )

    def _create_product_stock(self, product, quantity, unit_cost, extra_cost_total='0.00', purchase_date='2026-05-10'):
        purchase = Purchase.objects.create(
            organization=self.organization,
            purchase_date=purchase_date,
            counterparty=self.counterparty,
            created_by=self.owner,
            status=Purchase.Status.POSTED,
        )
        replace_purchase_items_and_stock(
            purchase,
            [
                {
                    'product': product,
                    'warehouse': self.warehouse,
                    'quantity': quantity,
                    'unit_cost': unit_cost,
                    'extra_cost_total': extra_cost_total,
                    'comment': '',
                }
            ],
            created_by=self.owner,
        )

    def test_products_page_filters_to_current_organization_and_supports_search(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.get(reverse('crm:products'))

        self.assertContains(response, 'SKU-001')
        self.assertContains(response, 'Базовый товар')
        self.assertNotContains(response, 'SKU-999')
        self.assertNotContains(response, 'Чужой товар')

        search_response = self.client.get(reverse('crm:products'), {'sku': '001', 'name': 'Базовый'})
        self.assertContains(search_response, 'SKU-001')
        self.assertNotContains(search_response, 'SKU-999')

    def test_products_page_uses_pagination(self):
        self.client.login(username='crm-products-owner', password='password123')

        for index in range(30):
            Product.objects.create(
                organization=self.organization,
                sku=f'SKU-PAGE-{index:03d}',
                name=f'Товар для пагинации {index:03d}',
            )

        first_page_response = self.client.get(reverse('crm:products'))
        second_page_response = self.client.get(reverse('crm:products'), {'page': 2})

        self.assertEqual(first_page_response.context['page_obj'].paginator.num_pages, 2)
        self.assertEqual(first_page_response.context['page_obj'].number, 1)
        self.assertEqual(len(first_page_response.context['products']), 25)
        self.assertEqual(second_page_response.context['page_obj'].number, 2)
        self.assertEqual(len(second_page_response.context['products']), 6)

    def test_products_page_filters_by_stock_quantity(self):
        self.client.login(username='crm-products-owner', password='password123')
        low_stock_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-LOW',
            name='Маленький остаток',
        )
        high_stock_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-HIGH',
            name='Большой остаток',
        )
        self._create_product_stock(low_stock_product, quantity=1, unit_cost='100.00')
        self._create_product_stock(high_stock_product, quantity=8, unit_cost='200.00')

        response = self.client.get(
            reverse('crm:products'),
            {
                'stock_from': '2',
            },
        )

        self.assertContains(response, 'SKU-HIGH')
        self.assertNotContains(response, 'SKU-LOW')
        self.assertNotContains(response, 'SKU-001')

    def test_products_page_supports_sorting_by_stock_and_cost(self):
        self.client.login(username='crm-products-owner', password='password123')
        low_stock_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-LOW',
            name='Маленький остаток',
        )
        high_stock_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-HIGH',
            name='Большой остаток',
        )
        self._create_product_stock(low_stock_product, quantity=1, unit_cost='150.00')
        self._create_product_stock(high_stock_product, quantity=8, unit_cost='500.00')

        stock_response = self.client.get(reverse('crm:products'), {'sort': 'stock_desc'})
        stock_products = list(stock_response.context['products'])
        self.assertEqual(stock_products[0].sku, 'SKU-HIGH')

        cost_response = self.client.get(reverse('crm:products'), {'sort': 'cost_asc'})
        cost_products = [product for product in cost_response.context['products'] if product.total_stock_quantity > 0]
        self.assertEqual(cost_products[0].sku, 'SKU-LOW')

    def test_product_create_shows_only_listing_channels_and_creates_offers_with_photo(self):
        self.client.login(username='crm-products-owner', password='password123')

        page_response = self.client.get(reverse('crm:product_create'))
        self.assertContains(page_response, 'Ozon')
        self.assertContains(page_response, 'Avito')
        self.assertNotContains(page_response, 'Самовывоз')
        self.assertNotContains(page_response, 'Старый маркетплейс')
        self.assertNotContains(page_response, 'Чужой канал')

        response = self.client.post(
            reverse('crm:product_create'),
            {
                'sku': 'SKU-NEW',
                'name': 'Новый товар',
                'comment': 'Комментарий',
                'is_active': 'on',
                'photo': SimpleUploadedFile(
                    'product.gif',
                    TEST_GIF_BYTES,
                    content_type='image/gif',
                ),
                f'channel_{self.ozon.id}_enabled': 'on',
                f'channel_{self.ozon.id}_price': '9990',
                f'channel_{self.ozon.id}_comment': 'Цена для Ozon',
                f'channel_{self.avito.id}_enabled': 'on',
                f'channel_{self.avito.id}_price': '9500',
                f'channel_{self.avito.id}_comment': 'Цена для Avito',
                f'channel_{self.other_listing_channel.id}_enabled': 'on',
                f'channel_{self.other_listing_channel.id}_price': '1',
            },
            follow=False,
        )

        self.assertRedirects(response, reverse('crm:products'))
        product = Product.objects.get(sku='SKU-NEW')
        self.assertEqual(product.organization, self.organization)
        self.assertTrue(product.is_active)
        self.assertTrue(product.photo.name.startswith('crm/products/'))

        offers = ProductChannelOffer.objects.filter(product=product).order_by('channel__name')
        self.assertEqual(offers.count(), 2)
        self.assertCountEqual(offers.values_list('channel__name', flat=True), ['Ozon', 'Avito'])
        self.assertEqual(str(ProductChannelOffer.objects.get(product=product, channel=self.ozon).price), '9990.00')
        self.assertFalse(
            ProductChannelOffer.objects.filter(product=product, channel=self.other_listing_channel).exists()
        )

    def test_product_create_requires_price_for_enabled_channel(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.post(
            reverse('crm:product_create'),
            {
                'sku': 'SKU-ERR',
                'name': 'Товар без цены',
                'is_active': 'on',
                f'channel_{self.ozon.id}_enabled': 'on',
                f'channel_{self.ozon.id}_price': '',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Укажите цену для канала, если товар на нем размещается.')
        self.assertFalse(Product.objects.filter(sku='SKU-ERR').exists())

    def test_product_edit_updates_fields_and_offer_matrix(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.post(
            reverse('crm:product_edit', args=[self.product.id]),
            {
                'sku': 'SKU-001-UPDATED',
                'name': 'Обновленный товар',
                'comment': 'Новый комментарий',
                'is_active': 'on',
                f'channel_{self.ozon.id}_price': '',
                f'channel_{self.ozon.id}_comment': 'Отключили Ozon',
                f'channel_{self.avito.id}_enabled': 'on',
                f'channel_{self.avito.id}_price': '9300',
                f'channel_{self.avito.id}_comment': 'Включили Avito',
            },
        )

        self.assertRedirects(response, reverse('crm:products'))
        self.product.refresh_from_db()
        self.assertEqual(self.product.sku, 'SKU-001-UPDATED')
        self.assertEqual(self.product.name, 'Обновленный товар')
        self.assertEqual(self.product.comment, 'Новый комментарий')

        ozon_offer = ProductChannelOffer.objects.get(product=self.product, channel=self.ozon)
        avito_offer = ProductChannelOffer.objects.get(product=self.product, channel=self.avito)
        self.assertFalse(ozon_offer.is_enabled)
        self.assertIsNone(ozon_offer.price)
        self.assertEqual(ozon_offer.comment, 'Отключили Ozon')
        self.assertTrue(avito_offer.is_enabled)
        self.assertEqual(str(avito_offer.price), '9300.00')
        self.assertEqual(avito_offer.comment, 'Включили Avito')

    def test_product_edit_page_loads_with_existing_photo(self):
        self.client.login(username='crm-products-owner', password='password123')
        self.product.photo = SimpleUploadedFile(
            'existing-product.gif',
            TEST_GIF_BYTES,
            content_type='image/gif',
        )
        self.product.save(update_fields=['photo', 'updated_at'])

        response = self.client.get(reverse('crm:product_edit', args=[self.product.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-photo-trigger')
        self.assertContains(response, 'data-photo-preview')
        self.assertNotContains(response, 'Расчетная себестоимость')

    def test_product_archive_sets_inactive_flag(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.post(reverse('crm:product_archive', args=[self.product.id]))

        self.assertRedirects(response, reverse('crm:products'))
        self.product.refresh_from_db()
        self.assertFalse(self.product.is_active)

    def test_user_cannot_open_or_change_other_organization_product(self):
        self.client.login(username='crm-products-owner', password='password123')

        edit_page_response = self.client.get(reverse('crm:product_edit', args=[self.other_product.id]))
        self.assertEqual(edit_page_response.status_code, 404)

        edit_response = self.client.post(
            reverse('crm:product_edit', args=[self.other_product.id]),
            {
                'sku': 'HACK',
                'name': 'Попытка',
            },
        )
        self.assertEqual(edit_response.status_code, 404)

        archive_response = self.client.post(reverse('crm:product_archive', args=[self.other_product.id]))
        self.assertEqual(archive_response.status_code, 404)

    def test_product_sku_is_unique_within_organization(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.post(
            reverse('crm:product_create'),
            {
                'sku': 'SKU-001',
                'name': 'Дубликат',
                'is_active': 'on',
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Товар с таким SKU / артикулом уже существует в этой организации.')
        self.assertEqual(Product.objects.filter(organization=self.organization, sku='SKU-001').count(), 1)

    def test_product_template_download_returns_expected_headers(self):
        self.client.login(username='crm-products-owner', password='password123')

        response = self.client.get(reverse('crm:product_template_download'))

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response['Content-Type'],
        )
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        sample_row = [cell.value for cell in worksheet[2]]
        self.assertEqual(headers, PRODUCT_IMPORT_TEMPLATE_HEADERS)
        self.assertEqual(sample_row, ['SKU-001', 'Пример товара', 'Короткий комментарий', 'Вставьте фото в эту строку'])

    def test_product_bulk_import_creates_only_new_products_groups_skipped_rows_and_imports_row_photos(self):
        self.client.login(username='crm-products-owner', password='password123')
        workbook_content = build_product_import_workbook(
            headers=['Номер', 'Набор', 'Комментарий', 'Фото'],
            rows=[
                ['SKU-001', 'Попытка обновления', 'Нельзя обновлять', ''],
                ['SKU-NEW-WITH-PHOTO', 'Новый товар с фото', 'Комментарий', ''],
                ['SKU-DUP', 'Первый дубль', 'Комментарий', ''],
                ['SKU-DUP', 'Второй дубль', 'Комментарий', ''],
                ['SKU-NEW-NO-PHOTO', 'Новый товар без фото', 'Комментарий', ''],
            ],
            image_cells=['D3'],
        )

        response = self.client.post(
            reverse('crm:product_bulk_import'),
            {
                'xlsx_file': SimpleUploadedFile(
                    'products.xlsx',
                    workbook_content,
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                ),
            },
            follow=True,
        )

        self.assertRedirects(response, reverse('crm:products'))
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, 'Базовый товар')
        self.assertEqual(self.product.comment, 'Тестовый комментарий')
        self.assertTrue(self.product.is_active)
        self.assertContains(response, 'Импорт завершен: загружено 3.')
        self.assertContains(response, 'Не загружено 2 строк.')
        self.assertContains(response, 'SKU-001')
        self.assertContains(response, 'SKU-DUP')
        self.assertContains(response, 'Уже существует в CRM')
        self.assertContains(response, 'Дублируется в файле')
        self.assertContains(response, 'Уже существует в CRM')
        self.assertContains(response, 'строки 2')
        self.assertContains(response, 'строки 5')

        imported_with_photo = Product.objects.get(organization=self.organization, sku='SKU-NEW-WITH-PHOTO')
        self.assertEqual(imported_with_photo.name, 'Новый товар с фото')
        self.assertTrue(imported_with_photo.is_active)
        self.assertTrue(imported_with_photo.photo.name.endswith('.png'))

        imported_duplicate = Product.objects.get(organization=self.organization, sku='SKU-DUP')
        self.assertEqual(imported_duplicate.name, 'Первый дубль')

        imported_without_photo = Product.objects.get(organization=self.organization, sku='SKU-NEW-NO-PHOTO')
        self.assertEqual(imported_without_photo.name, 'Новый товар без фото')
        self.assertTrue(imported_without_photo.is_active)
        self.assertFalse(imported_without_photo.photo)

    def test_product_bulk_import_requires_sku_and_name_headers(self):
        self.client.login(username='crm-products-owner', password='password123')
        workbook_content = build_product_import_workbook(
            headers=['Комментарий', 'Активен'],
            rows=[
                ['Комментарий без товара', '1'],
            ],
        )

        response = self.client.post(
            reverse('crm:product_bulk_import'),
            {
                'xlsx_file': SimpleUploadedFile(
                    'products.xlsx',
                    workbook_content,
                    content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'В Excel не хватает обязательных колонок')
        self.assertFalse(Product.objects.filter(organization=self.organization, sku='SKU-NEW').exists())


class CRMPurchasesTest(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='crm-purchases-owner', password='password123')
        self.other_owner = User.objects.create_user(username='crm-purchases-other', password='password123')

        self.organization = Organization.objects.create(
            owner=self.owner,
            name='Brick Trade',
            base_currency='RUB',
        )
        self.other_organization = Organization.objects.create(
            owner=self.other_owner,
            name='Other Org',
            base_currency='RUB',
        )

        Membership.objects.create(
            organization=self.organization,
            user=self.owner,
            role=Membership.Role.OWNER,
        )
        Membership.objects.create(
            organization=self.other_organization,
            user=self.other_owner,
            role=Membership.Role.OWNER,
        )

        self.main_warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Основной склад',
        )
        self.ozon_warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Ozon',
        )
        self.other_warehouse = Warehouse.objects.create(
            organization=self.other_organization,
            name='Чужой склад',
        )

        self.product = Product.objects.create(
            organization=self.organization,
            sku='SKU-001',
            name='Товар 1',
        )
        self.second_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-002',
            name='Товар 2',
        )
        self.other_product = Product.objects.create(
            organization=self.other_organization,
            sku='SKU-999',
            name='Чужой товар',
        )
        self.counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Поставщик X',
        )
        self.other_counterparty = Counterparty.objects.create(
            organization=self.other_organization,
            name='Чужой поставщик',
        )

    def _build_purchase_payload(self, rows, **purchase_fields):
        payload = {
            'purchase_date': purchase_fields.get('purchase_date', '2026-05-10'),
            'warehouse': str(purchase_fields.get('warehouse', self.main_warehouse).id),
            'counterparty': str(purchase_fields.get('counterparty', self.counterparty).id),
            'comment': purchase_fields.get('comment', 'Комментарий по закупке'),
            'status': purchase_fields.get('status', Purchase.Status.POSTED),
            'items-TOTAL_FORMS': str(len(rows)),
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
        }
        for index, row in enumerate(rows):
            payload[f'items-{index}-product'] = str(row['product'].id)
            payload[f'items-{index}-quantity'] = str(row['quantity'])
            payload[f'items-{index}-unit_cost'] = str(row['unit_cost'])
            payload[f'items-{index}-extra_cost_total'] = str(row.get('extra_cost_total', '0'))
            payload[f'items-{index}-comment'] = row.get('comment', '')
            payload[f'items-{index}-DELETE'] = ''
        return payload

    def _create_purchase(self):
        purchase = Purchase.objects.create(
            organization=self.organization,
            purchase_date='2026-05-01',
            warehouse=self.main_warehouse,
            counterparty=self.counterparty,
            comment='Первая закупка',
            created_by=self.owner,
            status=Purchase.Status.POSTED,
        )
        replace_purchase_items_and_stock(
            purchase,
            [
                {
                    'product': self.product,
                    'warehouse': self.main_warehouse,
                    'quantity': 2,
                    'unit_cost': '7000.00',
                    'extra_cost_total': '100.00',
                    'comment': 'Партия 1',
                }
            ],
            created_by=self.owner,
        )
        return purchase

    def test_purchase_create_creates_items_stock_lots_and_movements(self):
        self.client.login(username='crm-purchases-owner', password='password123')

        response = self.client.post(
            reverse('crm:purchase_create'),
            self._build_purchase_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 2,
                        'unit_cost': '7000.00',
                        'extra_cost_total': '100.00',
                        'comment': 'Партия 1',
                    },
                    {
                        'product': self.second_product,
                        'quantity': 1,
                        'unit_cost': '7200.00',
                        'extra_cost_total': '0.00',
                        'comment': 'Партия 2',
                    },
                ],
                counterparty=self.counterparty,
            ),
        )

        self.assertRedirects(response, reverse('crm:purchases'))
        purchase = Purchase.objects.get(counterparty=self.counterparty, purchase_date='2026-05-10')
        self.assertEqual(purchase.organization, self.organization)
        self.assertEqual(purchase.warehouse, self.main_warehouse)
        self.assertEqual(purchase.created_by, self.owner)
        self.assertEqual(purchase.status, Purchase.Status.POSTED)
        self.assertEqual(PurchaseItem.objects.filter(purchase=purchase).count(), 2)
        self.assertEqual(StockLot.objects.filter(organization=self.organization).count(), 2)
        self.assertEqual(
            StockMovement.objects.filter(
                organization=self.organization,
                movement_type=StockMovement.MovementType.PURCHASE_IN,
            ).count(),
            2,
        )

        lot = StockLot.objects.get(product=self.product)
        self.assertEqual(lot.warehouse, self.main_warehouse)
        self.assertEqual(lot.initial_quantity, 2)
        self.assertEqual(lot.remaining_quantity, 2)
        self.assertEqual(str(lot.unit_cost), '7000.00')
        self.assertEqual(str(lot.extra_cost_per_unit), '50.00')

    def test_purchases_page_and_template_download_work(self):
        purchase = self._create_purchase()
        self.client.login(username='crm-purchases-owner', password='password123')

        response = self.client.get(reverse('crm:purchases'))

        self.assertContains(response, 'Поставщик X')
        self.assertContains(response, 'Основной склад')
        self.assertContains(response, '2 шт.')
        self.assertContains(response, '14100.00')
        self.assertContains(response, 'Проведена')
        self.assertNotContains(response, 'Скачать шаблон CSV')

        template_response = self.client.get(reverse('crm:purchase_template_download'))
        self.assertEqual(template_response.status_code, 200)
        self.assertIn('text/csv', template_response['Content-Type'])
        content = template_response.content.decode('utf-8-sig')
        self.assertIn(','.join(PURCHASE_TEMPLATE_HEADERS), content)
        self.assertIn('SKU-001', content)

        detail_response = self.client.get(reverse('crm:purchase_detail', args=[purchase.id]))
        self.assertContains(detail_response, 'Партия 1')
        self.assertContains(detail_response, '50.00')

    def test_csv_upload_prefills_purchase_rows(self):
        self.client.login(username='crm-purchases-owner', password='password123')
        csv_content = (
            'sku,quantity,unit_cost,extra_cost_total,comment\n'
            'SKU-001,3,"7100,50","90,25",CSV строка\n'
            'SKU-002,1,"7200,00","0,00",CSV строка 2\n'
        ).encode('utf-8')

        response = self.client.post(
            reverse('crm:purchase_create'),
            {
                'purchase_date': '2026-05-10',
                'warehouse': str(self.main_warehouse.id),
                'counterparty': str(self.counterparty.id),
                'comment': 'CSV комментарий',
                'status': Purchase.Status.POSTED,
                'load_csv': '1',
                'file': SimpleUploadedFile('purchase.csv', csv_content, content_type='text/csv'),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(Purchase.objects.filter(organization=self.organization).count(), 0)
        self.assertEqual(response.context['csv_preview']['total_rows'], 2)
        self.assertEqual(response.context['item_formset'].forms[0].initial['quantity'], 3)
        self.assertEqual(str(response.context['item_formset'].forms[0].initial['unit_cost']), '7100.50')
        self.assertEqual(str(response.context['item_formset'].forms[0].initial['extra_cost_total']), '90.25')
        self.assertEqual(response.context['item_formset'].forms[1].initial['comment'], 'CSV строка 2')
        self.assertContains(response, 'CSV комментарий')
        self.assertContains(response, 'Предпросмотр CSV')
        self.assertContains(response, 'закупка еще не создана')

    def test_purchase_csv_upload_rejects_unknown_product_sku(self):
        self.client.login(username='crm-purchases-owner', password='password123')
        csv_content = (
            'sku,quantity,unit_cost,extra_cost_total,comment\n'
            'SKU-404,3,7100,90,Неизвестный товар\n'
        ).encode('utf-8')

        response = self.client.post(
            reverse('crm:purchase_create'),
            {
                'purchase_date': '2026-05-10',
                'warehouse': str(self.main_warehouse.id),
                'counterparty': str(self.counterparty.id),
                'comment': 'CSV комментарий',
                'status': Purchase.Status.POSTED,
                'load_csv': '1',
                'file': SimpleUploadedFile('purchase.csv', csv_content, content_type='text/csv'),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'товар с SKU &quot;SKU-404&quot; не найден в текущей организации')
        self.assertEqual(Purchase.objects.filter(organization=self.organization).count(), 0)

    def test_purchase_edit_rebuilds_items_and_stock(self):
        purchase = self._create_purchase()
        self.client.login(username='crm-purchases-owner', password='password123')

        response = self.client.post(
            reverse('crm:purchase_edit', args=[purchase.id]),
            self._build_purchase_payload(
                [
                    {
                        'product': self.second_product,
                        'quantity': 5,
                        'unit_cost': '7600.00',
                        'extra_cost_total': '200.00',
                        'comment': 'Обновленная партия',
                    }
                ],
                warehouse=self.ozon_warehouse,
                counterparty=self.counterparty,
                comment='Обновили закупку',
            ),
        )

        self.assertRedirects(response, reverse('crm:purchases'))
        purchase.refresh_from_db()
        self.assertEqual(purchase.counterparty, self.counterparty)
        self.assertEqual(PurchaseItem.objects.filter(purchase=purchase).count(), 1)
        self.assertFalse(StockLot.objects.filter(product=self.product).exists())
        updated_lot = StockLot.objects.get(product=self.second_product)
        self.assertEqual(updated_lot.remaining_quantity, 5)
        self.assertEqual(str(updated_lot.extra_cost_per_unit), '40.00')

    def test_draft_purchase_does_not_create_stock_lots_or_movements(self):
        self.client.login(username='crm-purchases-owner', password='password123')

        response = self.client.post(
            reverse('crm:purchase_create'),
            self._build_purchase_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 4,
                        'unit_cost': '6500.00',
                        'extra_cost_total': '0.00',
                    }
                ],
                status=Purchase.Status.DRAFT,
            ),
        )

        self.assertRedirects(response, reverse('crm:purchases'))
        purchase = Purchase.objects.get(purchase_date='2026-05-10')
        self.assertEqual(purchase.status, Purchase.Status.DRAFT)
        self.assertFalse(StockLot.objects.filter(purchase_item__purchase=purchase).exists())
        self.assertFalse(StockMovement.objects.filter(related_purchase_item__purchase=purchase).exists())

    def test_cancelled_purchase_removes_stock_effect(self):
        purchase = self._create_purchase()
        self.client.login(username='crm-purchases-owner', password='password123')

        response = self.client.post(reverse('crm:purchase_archive', args=[purchase.id]))

        self.assertRedirects(response, reverse('crm:purchases'))
        purchase.refresh_from_db()
        self.assertEqual(purchase.status, Purchase.Status.CANCELLED)
        self.assertFalse(StockLot.objects.filter(purchase_item__purchase=purchase).exists())
        self.assertFalse(StockMovement.objects.filter(related_purchase_item__purchase=purchase).exists())

    def test_purchase_access_is_scoped_to_current_organization(self):
        purchase = self._create_purchase()
        other_purchase = Purchase.objects.create(
            organization=self.other_organization,
            purchase_date='2026-05-02',
            warehouse=self.other_warehouse,
            counterparty=self.other_counterparty,
            created_by=self.other_owner,
            status=Purchase.Status.POSTED,
        )
        replace_purchase_items_and_stock(
            other_purchase,
            [
                {
                    'product': self.other_product,
                    'warehouse': self.other_warehouse,
                    'quantity': 1,
                    'unit_cost': '100.00',
                    'extra_cost_total': '0.00',
                    'comment': '',
                }
            ],
            created_by=self.other_owner,
        )

        self.client.login(username='crm-purchases-owner', password='password123')

        create_response = self.client.post(
            reverse('crm:purchase_create'),
            self._build_purchase_payload(
                [
                    {
                        'product': self.other_product,
                        'quantity': 1,
                        'unit_cost': '100.00',
                    }
                ]
            ),
        )
        self.assertEqual(create_response.status_code, 200)
        self.assertContains(create_response, 'Select a valid choice')

        foreign_warehouse_response = self.client.post(
            reverse('crm:purchase_create'),
            self._build_purchase_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 1,
                        'unit_cost': '100.00',
                    }
                ],
                warehouse=self.other_warehouse,
            ),
        )
        self.assertEqual(foreign_warehouse_response.status_code, 200)
        self.assertContains(foreign_warehouse_response, 'Select a valid choice')

        detail_response = self.client.get(reverse('crm:purchase_detail', args=[other_purchase.id]))
        self.assertEqual(detail_response.status_code, 404)

        edit_response = self.client.get(reverse('crm:purchase_edit', args=[other_purchase.id]))
        self.assertEqual(edit_response.status_code, 404)

        archive_response = self.client.post(reverse('crm:purchase_archive', args=[other_purchase.id]))
        self.assertEqual(archive_response.status_code, 404)

        product_response = self.client.get(reverse('crm:product_edit', args=[self.product.id]))
        self.assertContains(product_response, 'Основной склад')
        self.assertContains(product_response, '01.05.2026')
        self.assertContains(product_response, 'Поставщик X')


class CRMOrdersTest(TestCase):
    def setUp(self):
        self.owner = User.objects.create_user(username='crm-orders-owner', password='password123')
        self.other_owner = User.objects.create_user(username='crm-orders-other', password='password123')

        self.organization = Organization.objects.create(
            owner=self.owner,
            name='Brick Orders',
            base_currency='RUB',
        )
        self.other_organization = Organization.objects.create(
            owner=self.other_owner,
            name='Other Orders',
            base_currency='RUB',
        )

        Membership.objects.create(
            organization=self.organization,
            user=self.owner,
            role=Membership.Role.OWNER,
        )
        Membership.objects.create(
            organization=self.other_organization,
            user=self.other_owner,
            role=Membership.Role.OWNER,
        )

        self.avito = SalesChannel.objects.create(
            organization=self.organization,
            name='Avito',
            is_listing_channel=True,
        )
        self.pickup_channel = SalesChannel.objects.create(
            organization=self.organization,
            name='Самовывоз',
            is_listing_channel=False,
        )
        self.other_channel = SalesChannel.objects.create(
            organization=self.other_organization,
            name='Чужой канал',
            is_listing_channel=True,
        )

        self.courier = DeliveryMethod.objects.create(
            organization=self.organization,
            name='Курьер',
            code='courier',
            sort_order=10,
        )
        self.pickup_method = DeliveryMethod.objects.create(
            organization=self.organization,
            name='Самовывоз',
            code='pickup',
            sort_order=20,
        )
        self.other_delivery_method = DeliveryMethod.objects.create(
            organization=self.other_organization,
            name='Чужая доставка',
            code='other-delivery',
        )

        self.warehouse = Warehouse.objects.create(
            organization=self.organization,
            name='Основной склад',
        )
        self.other_warehouse = Warehouse.objects.create(
            organization=self.other_organization,
            name='Чужой склад',
        )
        self.counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Поставщик заказа',
        )
        self.customer_counterparty = Counterparty.objects.create(
            organization=self.organization,
            name='Постоянный клиент',
        )

        self.product = Product.objects.create(
            organization=self.organization,
            sku='SKU-ORDER-1',
            name='Товар для FIFO',
        )
        self.second_product = Product.objects.create(
            organization=self.organization,
            sku='SKU-ORDER-2',
            name='Второй товар',
        )
        self.other_product = Product.objects.create(
            organization=self.other_organization,
            sku='SKU-OTHER',
            name='Чужой товар',
        )

        ProductChannelOffer.objects.create(
            product=self.product,
            channel=self.avito,
            is_enabled=True,
            price='9500.00',
        )
        ProductChannelOffer.objects.create(
            product=self.second_product,
            channel=self.avito,
            is_enabled=True,
            price='12000.00',
        )

        self._create_stock_batch(self.product, quantity=1, unit_cost='7000.00', purchase_date='2026-05-01')
        self._create_stock_batch(self.product, quantity=3, unit_cost='7600.00', purchase_date='2026-05-10')
        self._create_stock_batch(self.second_product, quantity=2, unit_cost='5000.00', purchase_date='2026-05-12')

    def _create_stock_batch(self, product, quantity, unit_cost, purchase_date):
        purchase = Purchase.objects.create(
            organization=self.organization,
            purchase_date=purchase_date,
            counterparty=self.counterparty,
            created_by=self.owner,
            status=Purchase.Status.POSTED,
        )
        replace_purchase_items_and_stock(
            purchase,
            [
                {
                    'product': product,
                    'warehouse': self.warehouse,
                    'quantity': quantity,
                    'unit_cost': unit_cost,
                    'extra_cost_total': '0.00',
                    'comment': '',
                }
            ],
            created_by=self.owner,
        )
        return purchase

    def _build_order_payload(self, rows, **order_fields):
        payload = {
            'external_order_number': order_fields.get('external_order_number', 'EXT-123'),
            'order_date': order_fields.get('order_date', '2026-05-20'),
            'sales_channel': str(order_fields.get('sales_channel', self.avito).id),
            'delivery_method': str(order_fields.get('delivery_method', self.courier).id),
            'customer_counterparty': str(order_fields['customer_counterparty'].id) if order_fields.get('customer_counterparty') else '',
            'customer_name': order_fields.get('customer_name', 'Иван'),
            'customer_phone': order_fields.get('customer_phone', '+79990000000'),
            'customer_contact': order_fields.get('customer_contact', '@client'),
            'customer_address': order_fields.get('customer_address', 'Москва'),
            'customer_comment': order_fields.get('customer_comment', 'Позвонить заранее'),
            'payment_status': order_fields.get('payment_status', Order.PaymentStatus.UNPAID),
            'paid_amount': str(order_fields.get('paid_amount', '0.00')),
            'payment_method': order_fields.get('payment_method', ''),
            'status': order_fields.get('status', Order.Status.NEW),
            'comment': order_fields.get('comment', 'Комментарий по заказу'),
            'items-TOTAL_FORMS': str(len(rows)),
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
        }
        for index, row in enumerate(rows):
            payload[f'items-{index}-product'] = str(row['product'].id)
            payload[f'items-{index}-quantity'] = str(row['quantity'])
            payload[f'items-{index}-sale_price'] = row.get('sale_price', '')
            payload[f'items-{index}-comment'] = row.get('comment', '')
            payload[f'items-{index}-DELETE'] = ''
        return payload

    def _create_order(self, **order_kwargs):
        order = Order.objects.create(
            organization=self.organization,
            order_date=order_kwargs.get('order_date', '2026-05-20'),
            sales_channel=order_kwargs.get('sales_channel', self.avito),
            delivery_method=order_kwargs.get('delivery_method', self.courier),
            customer_name='Иван',
            status=order_kwargs.get('status', Order.Status.NEW),
            created_by=self.owner,
        )
        replace_order_items_and_reservations(
            order,
            order_kwargs.get(
                'rows',
                [
                    {
                        'product': self.product,
                        'quantity': 2,
                        'sale_price': None,
                        'comment': '',
                    }
                ],
            ),
        )
        order.refresh_from_db()
        return order

    def test_order_create_generates_internal_number_pulls_price_and_reserves_fifo(self):
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.post(
            reverse('crm:order_create'),
            self._build_order_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 2,
                        'sale_price': '',
                        'comment': 'Первая строка заказа',
                    }
                ]
            ),
        )

        order = Order.objects.get(organization=self.organization)
        self.assertRedirects(response, reverse('crm:order_detail', args=[order.id]))
        self.assertTrue(order.internal_order_number.startswith('ORD-'))
        self.assertEqual(order.status, Order.Status.AWAITING_SHIPMENT)

        order_item = OrderItem.objects.get(order=order)
        self.assertEqual(str(order_item.sale_price), '9500.00')
        reservations = list(
            StockReservation.objects.filter(order=order, status=StockReservation.Status.ACTIVE)
            .select_related('stock_lot__purchase_item__purchase')
            .order_by('stock_lot__purchase_item__purchase__purchase_date', 'id')
        )
        self.assertEqual([reservation.quantity for reservation in reservations], [1, 1])
        self.assertEqual(
            [reservation.stock_lot.purchase_item.purchase.purchase_date.isoformat() for reservation in reservations],
            ['2026-05-01', '2026-05-10'],
        )

        detail_response = self.client.get(reverse('crm:order_detail', args=[order.id]))
        self.assertContains(detail_response, 'Полный резерв')
        self.assertContains(detail_response, '19000.00')
        self.assertContains(detail_response, '14600.00')
        self.assertContains(detail_response, '4400.00')

    def test_order_form_uses_product_picker_and_hides_legacy_customer_fields(self):
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.get(reverse('crm:order_create'))

        self.assertContains(response, 'crmProductPickerModal')
        self.assertContains(response, 'Клиент из базы')
        self.assertNotContains(response, 'Комментарий клиента')
        self.assertNotContains(response, 'Способ оплаты')
        self.assertNotContains(response, 'Адрес')

    def test_order_create_can_link_customer_from_base_and_fill_name(self):
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.post(
            reverse('crm:order_create'),
            self._build_order_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 1,
                        'sale_price': '',
                    }
                ],
                customer_counterparty=self.customer_counterparty,
                customer_name='',
            ),
        )

        order = Order.objects.get(organization=self.organization)
        self.assertRedirects(response, reverse('crm:order_detail', args=[order.id]))
        self.assertEqual(order.customer_counterparty, self.customer_counterparty)
        self.assertEqual(order.customer_name, 'Постоянный клиент')

    def test_order_create_with_insufficient_stock_keeps_new_and_partial_reservation(self):
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.post(
            reverse('crm:order_create'),
            self._build_order_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 5,
                        'sale_price': '',
                    }
                ]
            ),
            follow=True,
        )

        order = Order.objects.get(organization=self.organization)
        self.assertEqual(order.status, Order.Status.NEW)
        self.assertContains(response, 'Не удалось полностью зарезервировать остатки')
        self.assertEqual(
            StockReservation.objects.filter(order=order, status=StockReservation.Status.ACTIVE).aggregate(
                total=Sum('quantity')
            )['total'],
            4,
        )

    def test_order_edit_recalculates_fifo_reservations(self):
        order = self._create_order(
            rows=[
                {
                    'product': self.product,
                    'quantity': 1,
                    'sale_price': None,
                    'comment': '',
                }
            ]
        )
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.post(
            reverse('crm:order_edit', args=[order.id]),
            self._build_order_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 3,
                        'sale_price': '',
                        'comment': '',
                    }
                ],
                external_order_number='UPDATED-1',
                status=Order.Status.AWAITING_SHIPMENT,
            ),
        )

        self.assertRedirects(response, reverse('crm:order_detail', args=[order.id]))
        order.refresh_from_db()
        self.assertEqual(order.external_order_number, 'UPDATED-1')
        self.assertEqual(order.status, Order.Status.AWAITING_SHIPMENT)
        reservations = list(
            StockReservation.objects.filter(order=order, status=StockReservation.Status.ACTIVE)
            .select_related('stock_lot__purchase_item__purchase')
            .order_by('stock_lot__purchase_item__purchase__purchase_date', 'id')
        )
        self.assertEqual([reservation.quantity for reservation in reservations], [1, 2])

    def test_order_cancel_releases_active_reservations(self):
        order = self._create_order()
        self.client.login(username='crm-orders-owner', password='password123')

        response = self.client.post(reverse('crm:order_cancel', args=[order.id]))

        self.assertRedirects(response, reverse('crm:orders'))
        order.refresh_from_db()
        self.assertEqual(order.status, Order.Status.CANCELLED)
        self.assertFalse(StockReservation.objects.filter(order=order, status=StockReservation.Status.ACTIVE).exists())
        self.assertEqual(
            StockReservation.objects.filter(order=order, status=StockReservation.Status.RELEASED).count(),
            2,
        )

    def test_orders_access_is_scoped_to_current_organization(self):
        other_order = Order.objects.create(
            organization=self.other_organization,
            order_date='2026-05-21',
            sales_channel=self.other_channel,
            delivery_method=self.other_delivery_method,
            status=Order.Status.NEW,
            created_by=self.other_owner,
        )
        replace_order_items_and_reservations(
            other_order,
            [
                {
                    'product': self.other_product,
                    'quantity': 1,
                    'sale_price': '100.00',
                    'comment': '',
                }
            ],
        )

        self.client.login(username='crm-orders-owner', password='password123')

        detail_response = self.client.get(reverse('crm:order_detail', args=[other_order.id]))
        self.assertEqual(detail_response.status_code, 404)

        edit_response = self.client.get(reverse('crm:order_edit', args=[other_order.id]))
        self.assertEqual(edit_response.status_code, 404)

        cancel_response = self.client.post(reverse('crm:order_cancel', args=[other_order.id]))
        self.assertEqual(cancel_response.status_code, 404)

        create_response = self.client.post(
            reverse('crm:order_create'),
            self._build_order_payload(
                [
                    {
                        'product': self.other_product,
                        'quantity': 1,
                        'sale_price': '100.00',
                    }
                ]
            ),
        )
        self.assertEqual(create_response.status_code, 200)
        self.assertContains(create_response, 'Select a valid choice')

        foreign_channel_response = self.client.post(
            reverse('crm:order_create'),
            self._build_order_payload(
                [
                    {
                        'product': self.product,
                        'quantity': 1,
                        'sale_price': '100.00',
                    }
                ],
                sales_channel=self.other_channel,
            ),
        )
        self.assertEqual(foreign_channel_response.status_code, 200)
        self.assertContains(foreign_channel_response, 'Select a valid choice')
